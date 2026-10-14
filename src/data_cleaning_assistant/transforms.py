from __future__ import annotations

import math
import shutil
from copy import copy, deepcopy
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import load_workbook
import pandas as pd
import yaml

from .actions import (
    ActionDecision,
    ActionExecution,
    ActionPreview,
    ActionType,
    CellChange,
    ChangeSet,
    CleaningAction,
    CleaningPlan,
    RowSnapshot,
    ValueSnapshot,
)
from .inference import is_missing_like, is_null, parse_date, parse_number
from .mapping_safety import selected_mapping_conflicts
from .planner import file_sha256


def _snapshot(value: Any) -> ValueSnapshot:
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        try:
            value = value.item()
        except (TypeError, ValueError):
            pass

    if is_null(value):
        return ValueSnapshot(kind="null", value=None)
    if isinstance(value, bool):
        return ValueSnapshot(kind="bool", value=value)
    if isinstance(value, int):
        return ValueSnapshot(kind="int", value=value)
    if isinstance(value, float):
        serialized = value if math.isfinite(value) else str(value)
        return ValueSnapshot(kind="float", value=serialized)
    if isinstance(value, datetime):
        return ValueSnapshot(kind="datetime", value=value.isoformat())
    if isinstance(value, date):
        return ValueSnapshot(kind="date", value=value.isoformat())
    if isinstance(value, str):
        return ValueSnapshot(kind="str", value=value)
    return ValueSnapshot(kind="str", value=str(value))


def _restore(snapshot: ValueSnapshot) -> Any:
    if snapshot.kind == "null":
        return None
    if snapshot.kind == "bool":
        return bool(snapshot.value)
    if snapshot.kind == "int":
        return int(snapshot.value)
    if snapshot.kind == "float":
        if snapshot.value == "nan":
            return float("nan")
        if snapshot.value == "inf":
            return float("inf")
        if snapshot.value == "-inf":
            return float("-inf")
        return float(snapshot.value)
    if snapshot.kind == "datetime":
        return datetime.fromisoformat(str(snapshot.value))
    if snapshot.kind == "date":
        return date.fromisoformat(str(snapshot.value))
    return str(snapshot.value)


def _row_snapshot(table: str, frame: pd.DataFrame, row_position: int) -> RowSnapshot:
    row = frame.iloc[row_position]
    return RowSnapshot(
        table=table,
        row_position=row_position,
        values={str(column): _snapshot(row[column]) for column in frame.columns},
    )


def _canonical_number(value: float) -> int | float:
    return int(value) if float(value).is_integer() else float(value)


def _new_value(action: CleaningAction, value: Any) -> tuple[bool, Any]:
    if action.type == ActionType.NORMALIZE_CATEGORIES:
        mapping = {
            str(key): replacement
            for key, replacement in action.parameters.get("mapping", {}).items()
        }
        key = str(value)
        if key in mapping and value != mapping[key]:
            return True, mapping[key]
        return False, value

    if action.type == ActionType.NORMALIZE_MISSING:
        tokens = {str(item) for item in action.parameters.get("values", [])}
        if not is_null(value) and str(value) in tokens:
            return True, None
        return False, value

    if action.type == ActionType.TRIM_WHITESPACE:
        if isinstance(value, str):
            stripped = value.strip()
            if stripped != value:
                return True, stripped
        return False, value

    if action.type == ActionType.STANDARDIZE_DATES:
        if is_missing_like(value):
            return False, value
        order = action.parameters.get("numeric_order")
        numeric_order = (
            str(order)
            if order in {"month_first", "day_first"}
            else None
        )
        parsed = parse_date(value, numeric_order=numeric_order)
        if parsed is None:
            return False, value
        fmt = str(action.parameters.get("format", "%Y-%m-%d"))
        standardized = parsed.strftime(fmt)
        if value != standardized:
            return True, standardized
        return False, value

    if action.type == ActionType.COERCE_NUMERIC:
        if is_missing_like(value):
            return False, value
        parsed = parse_number(value)
        if parsed is None:
            return False, value
        number = _canonical_number(parsed)
        if value != number or type(value) is not type(number):
            return True, number
        return False, value

    raise ValueError(f"Unsupported cell action type: {action.type}")


def _duplicate_positions(frame: pd.DataFrame) -> list[int]:
    duplicated = frame.duplicated(keep="first").tolist()
    return [position for position, matched in enumerate(duplicated) if bool(matched)]


def preview_action(
    tables: dict[str, pd.DataFrame],
    action: CleaningAction,
    sample_limit: int = 10,
) -> ActionPreview:
    if action.table not in tables:
        raise KeyError(f"Unknown table: {action.table}")
    frame = tables[action.table]

    if action.type == ActionType.REMOVE_EXACT_DUPLICATES:
        positions = _duplicate_positions(frame)
        return ActionPreview(
            action_id=action.id,
            affected_count=len(positions),
            sample_rows=[
                _row_snapshot(action.table, frame, position)
                for position in positions[:sample_limit]
            ],
        )

    if action.column is None or action.column not in frame.columns:
        raise KeyError(f"Unknown column {action.column!r} in table {action.table!r}")

    affected = 0
    samples: list[CellChange] = []
    for row_index, value in frame[action.column].items():
        changed, new_value = _new_value(action, value)
        if not changed:
            continue
        affected += 1
        if len(samples) < sample_limit:
            samples.append(
                CellChange(
                    table=action.table,
                    row_index=int(row_index),
                    column=action.column,
                    old_value=_snapshot(value),
                    new_value=_snapshot(new_value),
                )
            )
    return ActionPreview(
        action_id=action.id,
        affected_count=affected,
        sample_changes=samples,
    )


def apply_action(
    tables: dict[str, pd.DataFrame],
    action: CleaningAction,
) -> ActionExecution:
    if action.table not in tables:
        raise KeyError(f"Unknown table: {action.table}")
    frame = tables[action.table]

    if action.type == ActionType.REMOVE_EXACT_DUPLICATES:
        positions = _duplicate_positions(frame)
        deleted_rows = [_row_snapshot(action.table, frame, position) for position in positions]
        if positions:
            keep_mask = [position not in set(positions) for position in range(len(frame))]
            tables[action.table] = frame.iloc[keep_mask].copy()
        return ActionExecution(
            action_id=action.id,
            action_type=action.type,
            table=action.table,
            column=None,
            affected_count=len(deleted_rows),
            deleted_rows=deleted_rows,
        )

    if action.column is None or action.column not in frame.columns:
        raise KeyError(f"Unknown column {action.column!r} in table {action.table!r}")

    changes: list[CellChange] = []
    for row_index, value in frame[action.column].items():
        changed, new_value = _new_value(action, value)
        if not changed:
            continue
        changes.append(
            CellChange(
                table=action.table,
                row_index=int(row_index),
                column=action.column,
                old_value=_snapshot(value),
                new_value=_snapshot(new_value),
            )
        )
        frame.at[row_index, action.column] = new_value

    return ActionExecution(
        action_id=action.id,
        action_type=action.type,
        table=action.table,
        column=action.column,
        affected_count=len(changes),
        changes=changes,
    )


def apply_plan(
    path: str | Path,
    tables: dict[str, pd.DataFrame],
    plan: CleaningPlan,
    *,
    approved_ids: set[str] | None = None,
    approve_all: bool = False,
) -> tuple[dict[str, pd.DataFrame], ChangeSet]:
    resolved = Path(path).expanduser().resolve()
    actual_hash = file_sha256(resolved)
    if actual_hash != plan.source_sha256:
        raise ValueError(
            "The input file has changed since this cleaning plan was created. "
            "Generate a new plan before applying transformations."
        )

    approved_ids = approved_ids or set()
    known_ids = {action.id for action in plan.actions}
    unknown = approved_ids - known_ids
    if unknown:
        raise ValueError(f"Unknown action id(s): {', '.join(sorted(unknown))}")

    selected: list[CleaningAction] = []
    for action in plan.actions:
        if action.decision == ActionDecision.REJECTED:
            continue
        approved = (
            approve_all
            or action.decision == ActionDecision.APPROVED
            or action.id in approved_ids
        )
        if approved:
            selected.append(action)

    conflicts = selected_mapping_conflicts(selected)
    if conflicts:
        details = "; ".join(
            f"{conflict.table}.{conflict.column or '(table)'}: "
            f"{conflict.source!r} -> {list(conflict.targets)!r}"
            for conflict in conflicts[:8]
        )
        raise ValueError(
            "Conflicting approved category mappings would send the same source "
            f"value to multiple targets. Resolve the review conflict first. {details}"
        )

    # Cell edits run before row deletion. This preserves original row positions in
    # the audit log and makes persisted save/reload/undo deterministic.
    selected.sort(key=lambda action: action.type == ActionType.REMOVE_EXACT_DUPLICATES)

    cleaned = {name: frame.copy(deep=True) for name, frame in tables.items()}
    executions = []
    for action in selected:
        execution = apply_action(cleaned, action)
        if execution.affected_count:
            executions.append(execution)

    return cleaned, ChangeSet(
        source_path=str(resolved),
        source_sha256=actual_hash,
        executions=executions,
        study_schema_sha256=plan.study_schema_sha256,
        review_session=plan.review_session,
    )


def _resolve_table_key(
    tables: dict[str, pd.DataFrame],
    expected_table: str,
) -> str:
    if expected_table in tables:
        return expected_table
    if len(tables) == 1:
        return next(iter(tables))
    raise KeyError(f"Unknown table during undo: {expected_table}")


def _restore_deleted_rows(
    frame: pd.DataFrame,
    deleted_rows: list[RowSnapshot],
) -> pd.DataFrame:
    records = frame.to_dict(orient="records")
    for snapshot in sorted(deleted_rows, key=lambda row: row.row_position):
        restored_row = {
            column: _restore(snapshot.values[column])
            for column in frame.columns
        }
        records.insert(snapshot.row_position, restored_row)
    restored = pd.DataFrame(records, columns=frame.columns)
    for column, dtype in frame.dtypes.items():
        try:
            restored[column] = restored[column].astype(dtype)
        except (TypeError, ValueError):
            pass
    return restored


def undo_changeset(
    tables: dict[str, pd.DataFrame],
    changeset: ChangeSet,
) -> dict[str, pd.DataFrame]:
    restored = {name: frame.copy(deep=True) for name, frame in tables.items()}

    for execution in reversed(changeset.executions):
        table_key = _resolve_table_key(restored, execution.table)
        frame = restored[table_key]

        if execution.deleted_rows:
            frame = _restore_deleted_rows(frame, execution.deleted_rows)
            restored[table_key] = frame

        for change in reversed(execution.changes):
            if change.column not in frame.columns:
                raise KeyError(
                    f"Unknown column during undo: {execution.table}.{change.column}"
                )
            frame.at[change.row_index, change.column] = _restore(change.old_value)

    return restored


def write_tables(tables: dict[str, pd.DataFrame], output_path: str | Path) -> Path:
    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    suffix = output.suffix.lower()

    if suffix in {".xlsx", ".xlsm"}:
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            for table_name, frame in tables.items():
                frame.to_excel(writer, sheet_name=table_name[:31], index=False)
        return output

    if len(tables) != 1:
        raise ValueError(
            "CSV and Parquet outputs support one table. Use .xlsx for multi-sheet datasets."
        )

    frame = next(iter(tables.values()))
    if suffix == ".csv":
        frame.to_csv(output, index=False)
        return output
    if suffix == ".parquet":
        frame.to_parquet(output, index=False)
        return output
    raise ValueError("Output must end in .csv, .xlsx, .xlsm, or .parquet")


def _values_equal(left: Any, right: Any) -> bool:
    if is_null(left) and is_null(right):
        return True
    try:
        return bool(left == right)
    except (TypeError, ValueError):
        return repr(left) == repr(right)


def _sheet_has_formulas(worksheet) -> bool:
    return any(
        cell.data_type == "f"
        for row in worksheet.iter_rows()
        for cell in row
    )


def _copy_row_style(worksheet, source_row: int, target_row: int) -> None:
    if source_row < 1 or source_row > worksheet.max_row:
        return
    for column in range(1, worksheet.max_column + 1):
        source_cell = worksheet.cell(row=source_row, column=column)
        target_cell = worksheet.cell(row=target_row, column=column)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
    source_dimension = worksheet.row_dimensions[source_row]
    target_dimension = worksheet.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden


def _write_excel_preserving_source(
    source_path: Path,
    original_tables: dict[str, pd.DataFrame],
    cleaned_tables: dict[str, pd.DataFrame],
    output: Path,
    inserted_rows: dict[str, list[RowSnapshot]] | None = None,
) -> Path:
    keep_vba = source_path.suffix.lower() == ".xlsm"

    # Build beside the destination and publish only after a complete save. This
    # prevents a failed formula-safety check or write error from leaving behind
    # a misleading file at the requested cleaned-output path.
    with NamedTemporaryFile(
        prefix=f".{output.stem}.",
        suffix=output.suffix,
        dir=output.parent,
        delete=False,
    ) as handle:
        working = Path(handle.name)

    try:
        shutil.copy2(source_path, working)
        workbook = load_workbook(
            working,
            keep_vba=keep_vba,
            data_only=False,
        )

        for table_name, original in original_tables.items():
            if (
                table_name not in cleaned_tables
                or table_name not in workbook.sheetnames
            ):
                raise ValueError(
                    f"Cannot preserve Excel sheet {table_name!r}: "
                    "sheet mapping changed."
                )

            cleaned = cleaned_tables[table_name]
            worksheet = workbook[table_name]
            sheet_insertions = (inserted_rows or {}).get(table_name, [])

            if sheet_insertions:
                if _sheet_has_formulas(worksheet):
                    raise ValueError(
                        f"Refusing row insertion in sheet {table_name!r} "
                        "because it contains formulas."
                    )
                for snapshot in sorted(
                    sheet_insertions,
                    key=lambda row: row.row_position,
                ):
                    excel_row = snapshot.row_position + 2
                    worksheet.insert_rows(excel_row, 1)
                    style_source = (
                        excel_row - 1
                        if excel_row > 2
                        else excel_row + 1
                    )
                    _copy_row_style(
                        worksheet,
                        style_source,
                        excel_row,
                    )

                # Row reinsertion is used by audit-based undo. The cleaned
                # dataframe is already authoritative, and formula-bearing sheets
                # were rejected above, so values can be rewritten while styles
                # remain attached to their cells.
                for output_row, (_, row) in enumerate(
                    cleaned.iterrows(),
                    start=2,
                ):
                    for column_number, column in enumerate(
                        cleaned.columns,
                        start=1,
                    ):
                        value = row[column]
                        worksheet.cell(
                            row=output_row,
                            column=column_number,
                        ).value = None if is_null(value) else value
                continue

            cleaned_index = set(cleaned.index)
            deleted_positions = [
                position
                for position, index in enumerate(original.index)
                if index not in cleaned_index
            ]

            if deleted_positions and _sheet_has_formulas(worksheet):
                raise ValueError(
                    f"Refusing row deletion in sheet {table_name!r} "
                    "because it contains formulas. Removing worksheet rows can "
                    "change formula references. Review the workbook or export "
                    "to CSV/Parquet for row deletion."
                )

            for position in sorted(deleted_positions, reverse=True):
                worksheet.delete_rows(position + 2, 1)

            for output_row, original_index in enumerate(
                cleaned.index,
                start=2,
            ):
                if original_index not in original.index:
                    raise ValueError(
                        "Excel preservation currently supports cell edits and "
                        "audited row reinsertion, not arbitrary inserted rows."
                    )
                for column_number, column in enumerate(
                    cleaned.columns,
                    start=1,
                ):
                    desired = cleaned.at[original_index, column]
                    before = original.at[original_index, column]
                    if _values_equal(before, desired):
                        continue

                    cell = worksheet.cell(
                        row=output_row,
                        column=column_number,
                    )
                    if cell.data_type == "f":
                        raise ValueError(
                            "Refusing to overwrite formula cell "
                            f"{table_name}!{cell.coordinate}."
                        )
                    cell.value = None if is_null(desired) else desired

        workbook.save(working)
        working.replace(output)
        return output
    except Exception:
        working.unlink(missing_ok=True)
        raise


def write_cleaned_dataset(
    source_path: str | Path,
    original_tables: dict[str, pd.DataFrame],
    cleaned_tables: dict[str, pd.DataFrame],
    output_path: str | Path,
    *,
    inserted_rows: dict[str, list[RowSnapshot]] | None = None,
) -> Path:
    source = Path(source_path).expanduser().resolve()
    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    if source == output:
        raise ValueError("Refusing to overwrite the source dataset in place.")

    if (
        source.suffix.lower() in {".xlsx", ".xlsm"}
        and output.suffix.lower() in {".xlsx", ".xlsm"}
    ):
        return _write_excel_preserving_source(
            source,
            original_tables,
            cleaned_tables,
            output,
            inserted_rows=inserted_rows,
        )
    return write_tables(cleaned_tables, output)


def write_recipe(
    plan: CleaningPlan,
    changeset: ChangeSet,
    output_path: str | Path,
) -> Path:
    from .io import load_dataset
    from .recipe import build_schema_snapshot

    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    executed = {execution.action_id for execution in changeset.executions}
    actions = [
        {
            "type": action.type.value,
            "table": action.table,
            "column": action.column,
            "parameters": deepcopy(action.parameters),
        }
        for action in plan.actions
        if action.id in executed
    ]
    source_tables = load_dataset(changeset.source_path).tables
    payload = {
        "version": 2,
        "schema": build_schema_snapshot(source_tables),
        "actions": actions,
    }
    output.write_text(
        yaml.safe_dump(payload, sort_keys=False),
        encoding="utf-8",
    )
    return output
