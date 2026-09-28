"""Download and validate the public GRAPE clinical/VF workbook for workflow-pilot use.

Source:
  Huang X, Kong X, Shen Z, Jin K, Ye J.
  GRAPE dataset: VF and clinical information.xlsx
  Figshare article 23674350, DOI 10.6084/m9.figshare.23674350
  CC0

The downloader writes only into a fresh local output directory. It never modifies
another dataset and does not send the workbook to any model.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from datetime import UTC, datetime
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ARTICLE_URL = (
    "https://springernature.figshare.com/articles/dataset/"
    "GRAPE_dataset_VF_and_clinical_information_xlsx/23674350"
)
DOWNLOAD_URL = "https://springernature.figshare.com/ndownloader/files/41670009"
DOI = "10.6084/m9.figshare.23674350"
LICENSE = "CC0"
WORKBOOK = "grape_vf_clinical_information.xlsx"
MANIFEST = "source_manifest.json"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_workbook(path: Path) -> dict[str, int]:
    """Validate the published workbook's expected two-sheet shape."""
    book = load_workbook(path, read_only=True, data_only=False)
    expected = {"Baseline", "Follow-up"}
    if not expected.issubset(book.sheetnames):
        raise ValueError(
            f"Unexpected GRAPE workbook sheets: {book.sheetnames}; "
            f"required: {sorted(expected)}"
        )

    baseline = book["Baseline"]
    followup = book["Follow-up"]
    # Published article metadata reports 263 and 1115 data records respectively.
    baseline_rows = max(baseline.max_row - 1, 0)
    followup_rows = max(followup.max_row - 1, 0)
    if baseline_rows != 263 or followup_rows != 1115:
        raise ValueError(
            "GRAPE workbook row counts differ from the published description: "
            f"Baseline={baseline_rows}, Follow-up={followup_rows}"
        )
    if baseline.max_column < 70 or followup.max_column < 60:
        raise ValueError(
            "GRAPE workbook column counts are unexpectedly small; "
            "refusing to treat the download as the intended source."
        )
    return {
        "Baseline": baseline_rows,
        "Follow-up": followup_rows,
    }


def fetch(out: Path) -> Path:
    out = out.expanduser().resolve()
    if out.exists():
        raise FileExistsError(
            f"Output directory already exists: {out}. "
            "Choose a fresh path to preserve source provenance."
        )
    out.mkdir(parents=True, exist_ok=False)
    destination = out / WORKBOOK

    request = Request(
        DOWNLOAD_URL,
        headers={
            "User-Agent": (
                "Mozilla/5.0 DataCleaningAssistant/0.8 "
                "(public-research-workflow-pilot)"
            )
        },
    )
    try:
        with urlopen(request, timeout=60) as response, destination.open("wb") as handle:
            content_type = response.headers.get("Content-Type", "")
            while True:
                chunk = response.read(1024 * 1024)
                if not chunk:
                    break
                handle.write(chunk)
    except Exception:
        # Do not leave a partial workbook that could later be mistaken for source data.
        if destination.exists():
            destination.unlink()
        try:
            out.rmdir()
        except OSError:
            pass
        raise

    # XLSX is a ZIP container and should begin with the ZIP magic bytes.
    with destination.open("rb") as handle:
        if handle.read(2) != b"PK":
            destination.unlink(missing_ok=True)
            raise ValueError("Downloaded file is not an XLSX/ZIP container.")

    rows = validate_workbook(destination)
    manifest = {
        "version": 1,
        "purpose": "Public GRAPE workbook source provenance for local workflow pilot",
        "article_url": ARTICLE_URL,
        "download_url": DOWNLOAD_URL,
        "doi": DOI,
        "license": LICENSE,
        "retrieved_at_utc": datetime.now(UTC).isoformat(),
        "filename": WORKBOOK,
        "sha256": _sha256(destination),
        "bytes": destination.stat().st_size,
        "validated_data_rows": rows,
        "notes": (
            "Published workbook contains Baseline and Follow-up sheets. "
            "Article metadata states '/' represents absence of corresponding "
            "CFP or OCT examination in a visit. No model call or data edit occurred."
        ),
    }
    (out / MANIFEST).write_text(
        json.dumps(manifest, indent=2) + "\n",
        encoding="utf-8",
    )

    print("GRAPE public workbook downloaded and validated.")
    print(f"Baseline records: {rows['Baseline']}")
    print(f"Follow-up records: {rows['Follow-up']}")
    print(f"SHA-256: {manifest['sha256']}")
    print(f"Workbook: {destination}")
    print(f"Manifest: {out / MANIFEST}")
    return destination


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out",
        type=Path,
        default=Path("pilot_outputs/grape_public"),
        help="Fresh output directory for the public workbook and manifest",
    )
    args = parser.parse_args()
    fetch(args.out)


if __name__ == "__main__":
    main()
