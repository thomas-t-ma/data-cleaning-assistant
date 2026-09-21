"""Reproducible, local-only corruption benchmark for the public CoMMpass pilot.

The source workbook is NEVER edited. The extra BenchmarkLabels sheet contains
only synthetic test labels; it does not assert any clinical equivalence.
The answer key stores locations and non-identifying categorical values locally.

Run:
    python examples/commpass_challenge.py make --source PATH --schema YAML --out DIR
    data-cleaner profile DIR/commpass_challenge.xlsx --schema DIR/study_schema.yaml ...
    data-cleaner plan DIR/commpass_challenge.xlsx --schema DIR/study_schema.yaml ...
    python examples/commpass_challenge.py verify --challenge DIR --plan DIR/plan.json
    python examples/commpass_challenge.py verify --challenge DIR --plan DIR/plan.json \
        --exercise-apply-undo
"""
from __future__ import annotations

import argparse
import hashlib
import json
import random
import shutil
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from data_cleaning_assistant.actions import ActionType, CleaningPlan
from data_cleaning_assistant.ai_candidates import build_ai_candidates
from data_cleaning_assistant.engine import profile_dataset
from data_cleaning_assistant.io import load_dataset
from data_cleaning_assistant.planner import file_sha256
from data_cleaning_assistant.study_schema import load_study_schema
from data_cleaning_assistant.transforms import (
    apply_plan,
    preview_action,
    undo_changeset,
    write_cleaned_dataset,
)

WORKBOOK = "commpass_challenge.xlsx"
MANIFEST = "answer_key.json"
SCHEMA = "study_schema.yaml"
SYNTHETIC = "BenchmarkLabels"
SOURCE_SHEETS = ("Cases", "Diagnoses", "Treatments", "FollowUps", "LabTests")


def _columns(sheet: Any) -> dict[str, int]:
    return {
        str(cell.value): cell.column
        for cell in sheet[1]
        if cell.value is not None
    }


def _choose_blank(sheet: Any, column: str, rng: random.Random) -> int:
    index = _columns(sheet).get(column)
    if index is None:
        raise ValueError(f"Required pilot field absent: {sheet.title}.{column}")
    blank = [
        row for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, index).value in (None, "")
    ]
    present = [
        row for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, index).value not in (None, "")
    ]
    if not blank or not present:
        raise ValueError(
            f"Need both blank and observed entries for {sheet.title}.{column}"
        )
    return rng.choice(blank)


def _choose_category(
    sheet: Any, column: str, rng: random.Random, *, case_variant: bool
) -> tuple[int, str, str]:
    index = _columns(sheet).get(column)
    if index is None:
        raise ValueError(f"Required pilot field absent: {sheet.title}.{column}")
    rows: list[tuple[int, str]] = [
        (row, sheet.cell(row, index).value)
        for row in range(2, sheet.max_row + 1)
        if isinstance(sheet.cell(row, index).value, str)
        and sheet.cell(row, index).value.strip()
        and sheet.cell(row, index).data_type != "f"
    ]
    counts = Counter(value for _, value in rows)
    candidates: list[tuple[int, str, str]] = []
    for row, original in rows:
        if case_variant:
            transformed = original.upper() if original.upper() != original else original.lower()
            # The original must be the dominant category, not an existing
            # unrelated category, to support a deterministic canonical mapping.
            if (
                transformed == original
                or transformed in counts
                or counts[original] < 3
                or not any(char.isalpha() for char in original)
            ):
                continue
        else:
            transformed = f" {original} "
        candidates.append((row, original, transformed))
    if not candidates:
        raise ValueError(f"No safe injected category target: {sheet.title}.{column}")
    return rng.choice(candidates)


def _put(
    book: Any, table: str, column: str, excel_row: int, value: Any,
    kind: str, injections: list[dict[str, Any]],
) -> None:
    sheet = book[table]
    cell = sheet.cell(excel_row, _columns(sheet)[column])
    if cell.data_type == "f":
        raise ValueError(f"Refusing to replace a formula in {table}.{column}")
    before = cell.value
    if before == value:
        raise ValueError("Injection must actually change the cell")
    cell.value = value
    injections.append({
        "kind": kind,
        "table": table,
        "column": column,
        "row_index": excel_row - 2,
        "before": before,
        "after": value,
    })


def _make_labels(book: Any) -> None:
    if SYNTHETIC in book.sheetnames:
        raise ValueError(f"Source workbook already has reserved sheet {SYNTHETIC}")
    sheet = book.create_sheet(SYNTHETIC)
    sheet.append(["benchmark_row", "specimen_source", "assay_label"])
    samples = [
        ("Peripheral Blood", "Synthetic assay control 101"),
        ("Peripheral Blood", "Synthetic assay control 101"),
        ("Peripheral Blood", "Synthetic assay control 102"),
        ("Peripheral Blood", "Synthetic assay control 102"),
        ("Peripheral Blood", "Synthetic assay control 101"),
        ("Peripheral Blod", "Synthetic assay control 102"),
    ]
    for index, (source, assay) in enumerate(samples, start=1):
        sheet.append([f"B{index:03d}", source, assay])
    sheet.freeze_panes = "A2"


def make(source: Path, schema_path: Path, out: Path, seed: int) -> None:
    source = source.expanduser().resolve()
    schema_path = schema_path.expanduser().resolve()
    out = out.expanduser().resolve()
    if source.suffix.lower() != ".xlsx":
        raise ValueError("The public-data challenge expects an .xlsx pilot")
    if not source.is_file() or not schema_path.is_file():
        raise FileNotFoundError("Source workbook and original schema must exist")
    if out.exists():
        raise FileExistsError("Challenge output directory already exists; choose a fresh path")
    if out == source.parent:
        raise ValueError("Challenge output cannot replace the public pilot directory")

    original_hash = file_sha256(source)
    schema = load_study_schema(schema_path)
    loaded = load_dataset(source)
    schema.validate_dataset(loaded.tables)
    book = load_workbook(source)
    if not set(SOURCE_SHEETS).issubset(book.sheetnames):
        raise ValueError("Source must contain the five CoMMpass pilot worksheets")

    rng = random.Random(seed)
    injections: list[dict[str, Any]] = []
    for table, column, token in (
        ("Treatments", "days_to_treatment_end", "N/A"),
        ("FollowUps", "ecog_performance_status", "NA"),
    ):
        row = _choose_blank(book[table], column, rng)
        _put(book, table, column, row, token, "nonempty_null_sentinel", injections)

    for table, column, is_case, kind in (
        ("Cases", "primary_site", False, "surrounding_whitespace"),
        ("Cases", "disease_type", True, "category_case_variant"),
    ):
        row, _original, replacement = _choose_category(
            book[table], column, rng, case_variant=is_case
        )
        _put(book, table, column, row, replacement, kind, injections)

    _make_labels(book)
    if file_sha256(source) != original_hash:
        raise AssertionError("The original workbook changed before challenge export")
    out.mkdir(parents=True, exist_ok=False)
    challenge = out / WORKBOOK
    book.save(challenge)
    shutil.copyfile(schema_path, out / SCHEMA)

    manifest = {
        "version": 1,
        "purpose": "Deterministic public-data cleaner regression fixture",
        "seed": seed,
        "source_path": str(source),
        "source_sha256": original_hash,
        "source_schema_sha256": file_sha256(schema_path),
        "challenge_sha256": file_sha256(challenge),
        "injections": injections,
        "synthetic_sheet": SYNTHETIC,
        "protected_distinct_labels": [
            "Synthetic assay control 101", "Synthetic assay control 102",
        ],
        "notes": (
            "Synthetic BenchmarkLabels records are not clinical observations. "
            "Original source and schema remain unchanged. "
            "Injections store no patient identifiers."
        ),
    }
    (out / MANIFEST).write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    if file_sha256(source) != original_hash:
        raise AssertionError("Original workbook modified; investigate")
    print("Challenge created: 4 injected cell defects and 1 synthetic label sheet.")
    print("Original public workbook unchanged; no patient-level values printed.")
    print(f"Challenge: {challenge}")
    print(f"Answer key: {out / MANIFEST}")


def _candidate_pairs(report: Any, tables: dict[str, Any], schema: Any) -> set[frozenset[str]]:
    candidates = build_ai_candidates(report, tables, schema=schema)
    pairs: set[frozenset[str]] = set()
    for candidate in candidates:
        if candidate.action.table != SYNTHETIC:
            continue
        mapping = candidate.action.parameters.get("mapping", {})
        if len(mapping) == 1:
            old, new = next(iter(mapping.items()))
            pairs.add(frozenset((str(old), str(new))))
    return pairs


def _approved_safe_ids(
    tables: dict[str, Any], plan: CleaningPlan, injected: list[dict[str, Any]]
) -> tuple[set[str], dict[tuple[str, str, int], int], list[str]]:
    expected = {
        (item["table"], item["column"], item["row_index"]): item
        for item in injected
    }
    allowed: dict[tuple[str, str, int], int] = {}
    approved: set[str] = set()
    unexpected: list[str] = []
    for action in plan.actions:
        # Never approve anything on the synthetic semantic-label sheet, or
        # on any field outside the four explicitly injected test cells.
        if action.column is None:
            unexpected.append(f"{action.table}.{action.type.value}")
            continue
        targets = {
            key: item
            for key, item in expected.items()
            if key[:2] == (action.table, action.column)
        }
        if not targets:
            unexpected.append(f"{action.table}.{action.column}.{action.type.value}")
            continue
        if action.type not in {
            ActionType.NORMALIZE_MISSING,
            ActionType.TRIM_WHITESPACE,
            ActionType.NORMALIZE_CATEGORIES,
        }:
            unexpected.append(f"{action.table}.{action.column}.{action.type.value}")
            continue
        preview = preview_action(
            tables, action, sample_limit=len(tables[action.table])
        )
        changes = preview.sample_changes
        if not changes or len(changes) != preview.affected_count:
            unexpected.append(f"{action.table}.{action.column}.empty_or_incomplete_preview")
            continue
        valid = True
        for change in changes:
            key = (change.table, change.column, change.row_index)
            injected_item = expected.get(key)
            if injected_item is None:
                valid = False
                break
            if change.old_value.value != injected_item["after"]:
                valid = False
                break
            expected_after = (
                None if injected_item["kind"] == "nonempty_null_sentinel"
                else injected_item["before"]
            )
            if change.new_value.value != expected_after:
                valid = False
                break
            if (
                injected_item["kind"] == "nonempty_null_sentinel"
                and action.type != ActionType.NORMALIZE_MISSING
            ) or (
                injected_item["kind"] == "category_case_variant"
                and action.type != ActionType.NORMALIZE_CATEGORIES
            ) or (
                injected_item["kind"] == "surrounding_whitespace"
                and action.type not in {
                    ActionType.TRIM_WHITESPACE, ActionType.NORMALIZE_CATEGORIES,
                }
            ):
                valid = False
                break
            allowed[key] = allowed.get(key, 0) + 1
        if valid:
            approved.add(action.id)
        else:
            unexpected.append(f"{action.table}.{action.column}.{action.type.value}")
    return approved, allowed, unexpected


def verify(out: Path, plan_path: Path, exercise: bool) -> None:
    out = out.expanduser().resolve()
    plan_path = plan_path.expanduser().resolve()
    manifest = json.loads((out / MANIFEST).read_text(encoding="utf-8"))
    if manifest["version"] != 1:
        raise ValueError("Unsupported challenge answer-key version")
    challenge = out / WORKBOOK
    source = Path(manifest["source_path"])
    if file_sha256(challenge) != manifest["challenge_sha256"]:
        raise ValueError("Challenge workbook changed after fixture creation")
    if file_sha256(source) != manifest["source_sha256"]:
        raise ValueError("Original public workbook changed after fixture creation")
    if file_sha256(out / SCHEMA) != manifest["source_schema_sha256"]:
        raise ValueError("Challenge schema changed after fixture creation")

    plan = CleaningPlan.model_validate_json(plan_path.read_text(encoding="utf-8"))
    if plan.source_sha256 != manifest["challenge_sha256"]:
        raise ValueError("Cleaning plan is not for the exact challenge workbook")
    schema = load_study_schema(out / SCHEMA)
    loaded = load_dataset(challenge)
    schema.validate_dataset(loaded.tables)
    report = profile_dataset(challenge, schema=schema)
    expected = manifest["injections"]
    approved, covered, unexpected = _approved_safe_ids(
        loaded.tables, plan, expected
    )
    missing = sorted(
        f"{item['table']}.{item['column']}"
        for item in expected
        if (item["table"], item["column"], item["row_index"]) not in covered
    )
    pairs = _candidate_pairs(report, loaded.tables, schema)
    typo = frozenset(("Peripheral Blood", "Peripheral Blod"))
    distinct = frozenset(manifest["protected_distinct_labels"])
    # Detectors only suggest fuzzy label pairs, they never create
    # deterministic merge actions. Actual AI decisions are a separate test.
    fuzzy_flags = typo in pairs and distinct in pairs
    summaries = {
        "source_unchanged": True,
        "challenge_unchanged": True,
        "expected_injected_cells": len(expected),
        "safe_actions_eligible_for_fixture_only": len(approved),
        "expected_cells_covered": len(expected) - len(missing),
        "missing_expected_targets": missing,
        "unexpected_action_targets": unexpected,
        "synthetic_label_pairs_exposed_for_review": fuzzy_flags,
        "ambiguous_assay_labels_not_in_deterministic_merges": not any(
            action.table == SYNTHETIC for action in plan.actions
        ),
        "exercise_apply_undo": exercise,
    }
    if missing or unexpected or not fuzzy_flags:
        (out / "benchmark_score.json").write_text(
            json.dumps(summaries, indent=2) + "\n", encoding="utf-8"
        )
        raise AssertionError(
            "Benchmark proposal checks failed; see benchmark_score.json. "
            "No actions were applied."
        )

    if exercise:
        # Explicit opt-in applies ONLY ground-truth-validated synthetic
        # corruption fixes. Unreviewed proposals never receive approval.
        cleaned, audit = apply_plan(
            challenge, loaded.tables, plan, approved_ids=approved
        )
        if len(audit.executions) < 4:
            raise AssertionError("Fewer expected fixture changes executed")
        cleaned_path = out / "challenge_cleaned.xlsx"
        restored_path = out / "challenge_restored.xlsx"
        if cleaned_path.exists() or restored_path.exists():
            raise FileExistsError("Exercise outputs already exist; choose a new challenge dir")
        write_cleaned_dataset(challenge, loaded.tables, cleaned, cleaned_path)
        reloaded = load_dataset(cleaned_path)
        for item in expected:
            value = reloaded.tables[item["table"]].at[
                item["row_index"], item["column"]
            ]
            expected_value = (
                "" if item["kind"] == "nonempty_null_sentinel"
                else item["before"]
            )
            if value != expected_value:
                raise AssertionError("Cleaned cell does not match expected safe repair")

        restored = undo_changeset(reloaded.tables, audit)
        write_cleaned_dataset(cleaned_path, reloaded.tables, restored, restored_path)
        reread = load_dataset(restored_path)
        for name, original in loaded.tables.items():
            if not reread.tables[name].equals(original):
                raise AssertionError(f"Undo failed value-level equality on {name}")
        if file_sha256(challenge) != manifest["challenge_sha256"]:
            raise AssertionError("Challenge source changed during exercise")
        if file_sha256(source) != manifest["source_sha256"]:
            raise AssertionError("Original public source changed during exercise")
        audit_path = out / "challenge_audit.json"
        audit_path.write_text(audit.model_dump_json(indent=2), encoding="utf-8")
        summaries["applied_safe_cell_changes"] = audit.affected_cell_count
        summaries["undo_exact_value_equality"] = True
        summaries["source_unchanged"] = True
        summaries["challenge_unchanged"] = True
    (out / "benchmark_score.json").write_text(
        json.dumps(summaries, indent=2) + "\n", encoding="utf-8"
    )
    print(
        f"Benchmark passed: {summaries['expected_cells_covered']}/"
        f"{len(expected)} injected cells covered; "
        f"{len(approved)} reviewed fixture-safe actions; "
        "both synthetic fuzzy-label pairs exposed for semantic review."
    )
    if exercise:
        print("Fixture-only apply, workbook export, undo, and readback passed.")
        print("Original public workbook and challenge input remain unchanged.")
    print(f"Aggregate score: {out / 'benchmark_score.json'}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    maker = commands.add_parser("make", help="Make seeded corruption in a NEW directory")
    maker.add_argument("--source", type=Path, required=True)
    maker.add_argument("--schema", type=Path, required=True)
    maker.add_argument("--out", type=Path, required=True)
    maker.add_argument("--seed", type=int, default=20260921)
    checker = commands.add_parser("verify", help="Score plan against private answer key")
    checker.add_argument("--challenge", type=Path, required=True)
    checker.add_argument("--plan", type=Path, required=True)
    checker.add_argument(
        "--exercise-apply-undo", action="store_true",
        help="Explicit opt-in: exercise ONLY answer-key-validated fixture changes",
    )
    args = parser.parse_args()
    if args.command == "make":
        make(args.source, args.schema, args.out, args.seed)
    else:
        verify(args.challenge, args.plan, args.exercise_apply_undo)


if __name__ == "__main__":
    main()
