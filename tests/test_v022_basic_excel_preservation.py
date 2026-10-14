from __future__ import annotations

from copy import copy
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from data_cleaning_assistant.io import load_dataset
from data_cleaning_assistant.planner import propose_actions
from data_cleaning_assistant.transforms import apply_plan, write_cleaned_dataset


def _make_basic_workbook(path: Path) -> None:
    workbook = Workbook()
    employees = workbook.active
    employees.title = "Employees"
    employees.append(["Employee_ID", "Status", "Monthly_Pay", "Annualized"])
    employees.append(["E001", "Active", 5000, "=C2*12"])
    employees.append(["E002", "active", "N/A", '=IF(C3="","",C3*12)'])
    employees.append(["E003", " Active ", 4200, "=C4*12"])

    for cell in employees[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    employees["B2"].fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    employees.freeze_panes = "A2"
    employees.column_dimensions["A"].width = 18
    employees.column_dimensions["B"].width = 14

    lookup = workbook.create_sheet("Lookup")
    lookup.append(["Code", "Label"])
    lookup.append(["A", "Alpha"])
    lookup.append(["B", "Beta"])
    lookup["A1"].font = Font(bold=True)
    lookup.sheet_properties.tabColor = "00AA00"

    workbook.save(path)


def test_basic_excel_cleaning_preserves_formulas_styles_and_other_sheets(
    tmp_path: Path,
) -> None:
    source = tmp_path / "employees.xlsx"
    output = tmp_path / "employees.cleaned.xlsx"
    _make_basic_workbook(source)

    loaded = load_dataset(source)
    plan = propose_actions(source)
    cleaned, changeset = apply_plan(
        source,
        loaded.tables,
        plan,
        approve_all=True,
    )

    action_targets = {
        (action.type.value, action.table, action.column)
        for action in plan.actions
    }
    assert ("normalize_categories", "Employees", "Status") in action_targets
    assert ("normalize_missing", "Employees", "Monthly_Pay") in action_targets

    write_cleaned_dataset(
        source,
        loaded.tables,
        cleaned,
        output,
    )

    assert output.exists()
    assert len(changeset.executions) >= 2

    before = load_workbook(source, data_only=False)
    after = load_workbook(output, data_only=False)

    assert after.sheetnames == before.sheetnames == ["Employees", "Lookup"]

    employees = after["Employees"]
    assert employees["B2"].value == "Active"
    assert employees["B3"].value == "Active"
    assert employees["B4"].value == "Active"
    assert employees["C3"].value is None

    # Formula cells are not rewritten just because neighboring cells changed.
    assert employees["D2"].value == "=C2*12"
    assert employees["D3"].value == '=IF(C3="","",C3*12)'
    assert employees["D4"].value == "=C4*12"
    assert employees["D2"].data_type == "f"
    assert employees["D3"].data_type == "f"

    # Ordinary workbook presentation survives the local cell edits.
    assert employees["A1"].font.bold is True
    assert employees["A1"].fill.fgColor.rgb == before["Employees"]["A1"].fill.fgColor.rgb
    assert employees["B2"].fill.fgColor.rgb == before["Employees"]["B2"].fill.fgColor.rgb
    assert employees.freeze_panes == "A2"
    assert employees.column_dimensions["A"].width == before["Employees"].column_dimensions["A"].width

    # An unrelated sheet must be byte-for-value equivalent at the cell level.
    lookup_before = before["Lookup"]
    lookup_after = after["Lookup"]
    assert [
        [cell.value for cell in row]
        for row in lookup_after.iter_rows()
    ] == [
        [cell.value for cell in row]
        for row in lookup_before.iter_rows()
    ]
    assert lookup_after["A1"].font.bold is True
    assert lookup_after.sheet_properties.tabColor.rgb == lookup_before.sheet_properties.tabColor.rgb


def test_excel_duplicate_row_removal_fails_closed_when_sheet_has_formulas(
    tmp_path: Path,
) -> None:
    source = tmp_path / "formula_duplicates.xlsx"
    output = tmp_path / "formula_duplicates.cleaned.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["A", "B", "Computed"])
    sheet.append([1, 2, "=A2+B2"])
    sheet.append([1, 2, "=A3+B3"])
    workbook.save(source)

    loaded = load_dataset(source)
    plan = propose_actions(source)

    duplicate_actions = [
        action
        for action in plan.actions
        if action.type.value == "remove_exact_duplicates"
    ]
    assert len(duplicate_actions) == 1

    cleaned, _changeset = apply_plan(
        source,
        loaded.tables,
        plan,
        approve_all=True,
    )

    with pytest.raises(ValueError, match="because it contains formulas"):
        write_cleaned_dataset(
            source,
            loaded.tables,
            cleaned,
            output,
        )

    assert not output.exists() or output.stat().st_size > 0


def test_excel_writer_preserves_blank_and_duplicate_headers_when_other_cells_change(
    tmp_path: Path,
) -> None:
    source = tmp_path / "headers.xlsx"
    output = tmp_path / "headers.cleaned.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["ID", "Status", None, "Status"])
    sheet.append(["A001", " Active ", "free text", "Keep"])
    sheet.append(["A002", "Active", "more text", "Keep"])
    workbook.save(source)

    loaded = load_dataset(source)
    plan = propose_actions(source)
    cleaned, _changeset = apply_plan(
        source,
        loaded.tables,
        plan,
        approve_all=True,
    )
    write_cleaned_dataset(source, loaded.tables, cleaned, output)

    after = load_workbook(output, data_only=False)["Data"]
    assert [after.cell(row=1, column=i).value for i in range(1, 5)] == [
        "ID",
        "Status",
        None,
        "Status",
    ]
    assert after["B2"].value == "Active"
    assert after["C2"].value == "free text"
    assert after["D2"].value == "Keep"
