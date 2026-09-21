from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook
from typer.testing import CliRunner

from data_cleaning_assistant.actions import ActionDecision, ActionType, CleaningPlan
from data_cleaning_assistant.ai_planner import propose_ai_plan
from data_cleaning_assistant.cli import app
from data_cleaning_assistant.engine import profile_dataset
from data_cleaning_assistant.io import load_dataset
from data_cleaning_assistant.planner import propose_actions
from data_cleaning_assistant.study_schema import StudySchema, load_study_schema
from data_cleaning_assistant.transforms import apply_plan


class DeferProvider:
    name = "fake"

    def generate_structured(
        self,
        *,
        model: str,
        system_prompt: str,
        user_prompt: str,
        response_model,
    ):
        assert "C001" in user_prompt
        return response_model.model_validate(
            {
                "summary": "The category labels need dictionary confirmation.",
                "selections": [
                    {
                        "candidate_id": "C001",
                        "decision": "defer",
                        "rationale": "Check the study codebook before merging the labels.",
                        "confidence": 0.95,
                    }
                ],
            }
        )


def _write_longitudinal_workbook(path: Path) -> None:
    wb = Workbook()
    subjects = wb.active
    subjects.title = "Subjects"
    subjects.append(["SubjectID", "Status"])
    subjects.append(["S001", "Known"])
    subjects.append(["S002", "Known"])

    visits = wb.create_sheet("Visits")
    visits.append(["SubjectID", "VisitOrder", "VisitDate", "Arm", "Status"])
    for values in (
        ("S001", "1", "2026-02-03", "Armm A", "Known"),
        ("S001", "2", "2026-01-01", "Arm A", "Unknown"),
        ("S001", "2", "2026-01-02", "Arm B", "Known"),
        ("S003", "1", "2026-01-01", "Arm C", "Known"),
        ("S002", "1", "2026-02-01", "Arm B", "Known"),
    ):
        visits.append(values)
    wb.save(path)


def _schema_text() -> str:
    return (
        "version: 1\n"
        "tables:\n"
        "  Subjects:\n"
        "    grain_key: [SubjectID]\n"
        "  Visits:\n"
        "    grain_key: [SubjectID, VisitOrder]\n"
        "    subject_id: SubjectID\n"
        "    time_column: VisitDate\n"
        "    order_column: VisitOrder\n"
        "    columns:\n"
        "      Arm:\n"
        "        allowed_values: [Arm A, Arm B]\n"
        "        canonical_map:\n"
        "          Armm A: Arm A\n"
        "      Status:\n"
        "        missing_tokens: [Unknown]\n"
        "foreign_keys:\n"
        "  - child_table: Visits\n"
        "    child_column: SubjectID\n"
        "    parent_table: Subjects\n"
        "    parent_column: SubjectID\n"
    )


def test_schema_profiles_cross_table_and_longitudinal_constraints(tmp_path: Path) -> None:
    source = tmp_path / "study.xlsx"
    schema_path = tmp_path / "study.yaml"
    _write_longitudinal_workbook(source)
    schema_path.write_text(_schema_text(), encoding="utf-8")
    schema = load_study_schema(schema_path)

    report = profile_dataset(source, schema=schema)
    issue_types = {issue.type for issue in report.issues}
    assert "schema_duplicate_grain" in issue_types
    assert "schema_visit_chronology" in issue_types
    assert "schema_unmatched_foreign_key" in issue_types
    assert "schema_unexpected_category" in issue_types
    assert all(
        issue.auto_fix_safe is False
        for issue in report.issues
        if issue.type.startswith("schema_")
    )

    foreign = next(i for i in report.issues if i.type == "schema_unmatched_foreign_key")
    assert foreign.affected_count == 1
    assert "S003" not in str(foreign.model_dump())
    chronology = next(i for i in report.issues if i.type == "schema_visit_chronology")
    assert chronology.affected_count >= 1


def test_explicit_schema_authorizes_only_declared_edits(tmp_path: Path) -> None:
    source = tmp_path / "study.xlsx"
    schema_path = tmp_path / "study.yaml"
    _write_longitudinal_workbook(source)
    schema_path.write_text(_schema_text(), encoding="utf-8")
    schema = load_study_schema(schema_path)

    plan = propose_actions(source, schema=schema)
    assert plan.study_schema_sha256 == schema.source_sha256
    categories = [
        action for action in plan.actions
        if action.type == ActionType.NORMALIZE_CATEGORIES and action.column == "Arm"
    ]
    assert len(categories) == 1
    assert categories[0].parameters == {"mapping": {"Armm A": "Arm A"}}
    missing = [
        action for action in plan.actions
        if action.type == ActionType.NORMALIZE_MISSING and action.column == "Status"
    ]
    assert len(missing) == 1
    assert missing[0].parameters == {"values": ["Unknown"]}
    assert all(action.decision == ActionDecision.PROPOSED for action in plan.actions)

    loaded = load_dataset(source)
    cleaned, audit = apply_plan(
        source,
        loaded.tables,
        plan,
        approved_ids={categories[0].id, missing[0].id},
    )
    assert cleaned["Visits"]["Arm"].iloc[0] == "Arm A"
    assert pd.isna(cleaned["Visits"]["Status"].iloc[1])
    assert audit.study_schema_sha256 == schema.source_sha256


def test_no_schema_preserves_contextual_missingness(tmp_path: Path) -> None:
    source = tmp_path / "study.csv"
    source.write_text(
        "SubjectID,Status\nS001,Known\nS002,Unknown\nS003,Known\n",
        encoding="utf-8",
    )
    plan = propose_actions(source)
    assert not any(a.type == ActionType.NORMALIZE_MISSING for a in plan.actions)


def test_schema_rejects_unknown_columns_and_identifier_edits(tmp_path: Path) -> None:
    source = tmp_path / "study.csv"
    source.write_text("SubjectID,Arm\nS001,A\n", encoding="utf-8")
    frame = load_dataset(source).tables
    schema = StudySchema.model_validate(
        {"tables": {"study": {"columns": {"MissingColumn": {"allowed_values": ["A"]}}}}}
    )
    with pytest.raises(ValueError, match="unknown columns"):
        schema.validate_dataset(frame)

    schema = StudySchema.model_validate(
        {"tables": {
            "study": {
                "subject_id": "SubjectID",
                "columns": {"SubjectID": {"canonical_map": {"S001": "S002"}}},
            }
        }}
    )
    with pytest.raises(ValueError, match="identifier edits"):
        schema.validate_dataset(frame)


def test_schema_rejects_inconsistent_canonical_dictionary() -> None:
    with pytest.raises(ValueError, match="canonical_map targets"):
        StudySchema.model_validate(
            {"tables": {"Visits": {"columns": {
                "Arm": {
                    "allowed_values": ["Arm A"],
                    "canonical_map": {"Armm A": "Arm B"},
                }
            }}}}
        )


def test_ai_deferral_is_persisted_but_never_auto_applied(tmp_path: Path) -> None:
    source = tmp_path / "study.csv"
    source.write_text(
        "patient_id,group\n"
        "P001,Baseline\n"
        "P002,Basline\n"
        "P003,Baseline\n"
        "P004,Baseline\n",
        encoding="utf-8",
    )
    result = propose_ai_plan(
        source,
        provider=DeferProvider(),
        model="fake",
        include_rules=False,
    )
    assert result.plan.actions == []
    assert len(result.plan.deferred_reviews) == 1
    deferred = result.plan.deferred_reviews[0]
    assert deferred.rationale.startswith("Check the study codebook")
    assert deferred.proposed_action.parameters == {"mapping": {"Basline": "Baseline"}}
    assert deferred.status == "pending"

    loaded = load_dataset(source)
    unchanged, audit = apply_plan(source, loaded.tables, result.plan, approve_all=True)
    assert audit.executions == []
    assert unchanged["study"].equals(loaded.tables["study"])

    plan_file = tmp_path / "ai_plan.json"
    reviewed_file = tmp_path / "reviewed.json"
    plan_file.write_text(result.plan.model_dump_json(indent=2), encoding="utf-8")
    review = CliRunner().invoke(
        app,
        ["review", str(source), str(plan_file), "-o", str(reviewed_file)],
        input="a\n",
    )
    assert review.exit_code == 0, review.output
    reviewed = CleaningPlan.model_validate_json(reviewed_file.read_text(encoding="utf-8"))
    assert reviewed.deferred_reviews[0].status == "approved"
    assert len(reviewed.actions) == 1
    assert reviewed.actions[0].decision == ActionDecision.APPROVED
    cleaned, audit = apply_plan(source, loaded.tables, reviewed)
    assert audit.affected_cell_count == 1
    assert cleaned["study"]["group"].iloc[1] == "Baseline"


def test_cli_profile_plan_accept_schema(tmp_path: Path) -> None:
    source = tmp_path / "study.xlsx"
    schema_path = tmp_path / "study.yaml"
    report_path = tmp_path / "report.json"
    plan_path = tmp_path / "plan.json"
    _write_longitudinal_workbook(source)
    schema_path.write_text(_schema_text(), encoding="utf-8")

    runner = CliRunner()
    profile_result = runner.invoke(
        app, ["profile", str(source), "--schema", str(schema_path), "-o", str(report_path)]
    )
    assert profile_result.exit_code == 0, profile_result.output
    assert "schema_unmatched_foreign_key" in report_path.read_text(encoding="utf-8")

    plan_result = runner.invoke(
        app, ["plan", str(source), "--schema", str(schema_path), "-o", str(plan_path)]
    )
    assert plan_result.exit_code == 0, plan_result.output
    plan = CleaningPlan.model_validate_json(plan_path.read_text(encoding="utf-8"))
    assert plan.study_schema_sha256


def test_schema_canonical_mapping_does_not_need_ai_acceptance(tmp_path: Path) -> None:
    source = tmp_path / "study.csv"
    source.write_text(
        "specimen_id,specimen_source\n"
        "SP001,Peripheral Blood\n"
        "SP002,Peripheral Blood\n"
        "SP003,Peripheral Blod\n"
        "SP004,Peripheral Blood\n",
        encoding="utf-8",
    )
    schema = StudySchema.model_validate(
        {"tables": {
            "study": {
                "columns": {
                    "specimen_source": {
                        "allowed_values": ["Peripheral Blood", "Bone Marrow"],
                        "canonical_map": {"Peripheral Blod": "Peripheral Blood"},
                    }
                }
            }
        }}
    )

    class MustNotCallAI:
        name = "fake"

        def generate_structured(self, **kwargs):
            raise AssertionError("Schema-authorized mapping must not depend on AI judgment")

    result = propose_ai_plan(
        source,
        provider=MustNotCallAI(),
        model="fake",
        schema=schema,
    )
    assert result.candidates == []
    categories = [
        a for a in result.plan.actions
        if a.type == ActionType.NORMALIZE_CATEGORIES
        and a.column == "specimen_source"
    ]
    assert len(categories) == 1
    assert categories[0].parameters == {
        "mapping": {"Peripheral Blod": "Peripheral Blood"}
    }
    assert categories[0].decision == ActionDecision.PROPOSED
    assert result.plan.deferred_reviews == []
