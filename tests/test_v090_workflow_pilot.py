from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from data_cleaning_assistant.actions import ActionDecision
from data_cleaning_assistant.ai_models import CandidateDecision
from data_cleaning_assistant.planner import file_sha256


class PilotFakeProvider:
    name = "offline_fake"

    def __init__(self) -> None:
        self.calls = 0

    def generate_structured(
        self, *, model, system_prompt, user_prompt, response_model
    ):
        assert model == "fake-model"
        assert system_prompt
        self.calls += 1
        context = json.loads(user_prompt.split("\n\n", 1)[1])
        selections = []
        for candidate in context["candidates"]:
            decision = (
                CandidateDecision.ACCEPT
                if candidate["column"] == "specimen_source"
                else CandidateDecision.DEFER
            )
            selections.append(
                {
                    "candidate_id": candidate["id"],
                    "decision": decision.value,
                    "rationale": f"fixture {decision.value}",
                    "confidence": 0.9,
                }
            )
        return response_model.model_validate(
            {"summary": "offline pilot fixture", "selections": selections}
        )


def _source(path: Path) -> None:
    path.write_text(
        "subject_id,specimen_source,stage\n"
        "S001,Peripheral Blood,Stage II\n"
        "S002,Peripheral Blood,Stage III\n"
        "S003,Peripheral Blod,Stage II\n"
        "S004,Peripheral Blood,Stage III\n",
        encoding="utf-8",
    )


def test_workflow_pilot_start_and_score_human_disagreement(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[1] / "examples"))
    from workflow_pilot import (
        CONTEXT,
        MANIFEST,
        PLAN,
        PROFILE,
        SCORE,
        TRACE,
        score_pilot,
        start_pilot,
    )

    source = tmp_path / "research.csv"
    _source(source)
    before = file_sha256(source)
    pilot = tmp_path / "pilot"

    provider = PilotFakeProvider()
    started = start_pilot(
        source,
        pilot,
        provider=provider,
        model="fake-model",
    )

    assert provider.calls == 2
    assert file_sha256(source) == before
    assert all((pilot / name).exists() for name in (MANIFEST, PROFILE, CONTEXT, PLAN, TRACE))
    assert started["trace"]["candidate_count"] == 2
    assert {row["decision"] for row in started["trace"]["candidates"]} == {
        "accept",
        "defer",
    }

    reviewed = started["plan"].model_copy(deep=True)
    accepted_candidate = next(
        row for row in started["trace"]["candidates"] if row["decision"] == "accept"
    )
    assert accepted_candidate["covered_by_plan_action_ids"]
    accepted_action_id = accepted_candidate["covered_by_plan_action_ids"][0]
    next(a for a in reviewed.actions if a.id == accepted_action_id).decision = (
        ActionDecision.REJECTED
    )

    deferred_candidate = next(
        row for row in started["trace"]["candidates"] if row["decision"] == "defer"
    )
    deferred = next(
        item
        for item in reviewed.deferred_reviews
        if item.candidate_id == deferred_candidate["candidate_id"]
    )
    deferred.status = "approved"
    approved_action = deferred.proposed_action.model_copy(deep=True)
    approved_action.decision = ActionDecision.APPROVED
    reviewed.actions.append(approved_action)

    reviewed_path = pilot / "reviewed_plan.json"
    reviewed_path.write_text(reviewed.model_dump_json(indent=2), encoding="utf-8")

    scored = score_pilot(
        pilot,
        reviewed_path,
        review_minutes=4.5,
        missed_issues=["example reviewer-observed issue"],
    )

    assert scored["resolved_candidate_reviews"] == 2
    assert scored["human_ai_agreement_count"] == 0
    assert scored["human_rejected_ai_accept_count"] == 1
    assert scored["human_approved_ai_defer_count"] == 1
    assert scored["review_minutes"] == 4.5
    assert scored["missed_issues_reported_by_reviewer"] == [
        "example reviewer-observed issue"
    ]
    assert (pilot / SCORE).exists()
    assert file_sha256(source) == before


def test_workflow_pilot_refuses_changed_source(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[1] / "examples"))
    from workflow_pilot import score_pilot, start_pilot

    source = tmp_path / "research.csv"
    _source(source)
    pilot = tmp_path / "pilot"
    started = start_pilot(
        source,
        pilot,
        provider=PilotFakeProvider(),
        model="fake-model",
    )
    reviewed_path = pilot / "reviewed_plan.json"
    reviewed_path.write_text(
        started["plan"].model_dump_json(indent=2),
        encoding="utf-8",
    )

    source.write_text(source.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    with pytest.raises(ValueError, match="source is missing or changed"):
        score_pilot(pilot, reviewed_path)


def test_grape_fetcher_validates_expected_sheet_shape(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[1] / "examples"))
    from fetch_grape_workbook import validate_workbook

    workbook = tmp_path / "grape.xlsx"
    book = Workbook()
    baseline = book.active
    baseline.title = "Baseline"
    baseline.append([f"c{i}" for i in range(80)])
    for _ in range(263):
        baseline.append([None] * 80)

    followup = book.create_sheet("Follow-up")
    followup.append([f"c{i}" for i in range(69)])
    for _ in range(1115):
        followup.append([None] * 69)
    book.save(workbook)

    assert validate_workbook(workbook) == {
        "Baseline": 263,
        "Follow-up": 1115,
    }


def test_grape_fetcher_rejects_wrong_shape(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[1] / "examples"))
    from fetch_grape_workbook import validate_workbook

    workbook = tmp_path / "wrong.xlsx"
    book = Workbook()
    baseline = book.active
    baseline.title = "Baseline"
    baseline.append(["x"] * 80)
    baseline.append(["x"] * 80)
    followup = book.create_sheet("Follow-up")
    followup.append(["x"] * 69)
    followup.append(["x"] * 69)
    book.save(workbook)

    with pytest.raises(ValueError, match="row counts differ"):
        validate_workbook(workbook)
