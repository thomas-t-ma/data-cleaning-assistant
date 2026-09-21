"""Offline, seeded public-workbook challenge; no GDC/API access in CI."""

from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from data_cleaning_assistant.actions import CleaningPlan
from data_cleaning_assistant.planner import file_sha256

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "examples" / "commpass_challenge.py"


def _fixture(path: Path) -> None:
    book = Workbook()
    cases = book.active
    cases.title = "Cases"
    cases.append(["case_id", "disease_type", "primary_site"])
    for i in range(1, 6):
        cases.append([f"P{i:03d}", "Plasma Cell Myeloma", "Bone Marrow"])

    diagnoses = book.create_sheet("Diagnoses")
    diagnoses.append(["diagnosis_key", "case_id", "age_at_diagnosis"])
    for i in range(1, 6):
        diagnoses.append([f"P{i:03d}:D0001", f"P{i:03d}", 50 + i])

    treatments = book.create_sheet("Treatments")
    treatments.append([
        "treatment_key", "diagnosis_key", "case_id",
        "days_to_treatment_end", "treatment_type",
    ])
    for i in range(1, 6):
        treatments.append([
            f"P{i:03d}:D0001:T0001", f"P{i:03d}:D0001",
            f"P{i:03d}", i * 100 if i <= 3 else None, "Systemic",
        ])

    followups = book.create_sheet("FollowUps")
    followups.append(["followup_key", "case_id", "ecog_performance_status"])
    for i in range(1, 6):
        followups.append([f"P{i:03d}:F0001", f"P{i:03d}", 0 if i <= 3 else None])

    labs = book.create_sheet("LabTests")
    labs.append(["labtest_key", "followup_key", "case_id", "laboratory_test", "test_value"])
    for i in range(1, 6):
        labs.append([
            f"P{i:03d}:F0001:L0001", f"P{i:03d}:F0001", f"P{i:03d}",
            "Synthetic assay", i + 1,
        ])
    book.save(path)


def _call(*args: str) -> subprocess.CompletedProcess[str]:
    env = dict(os.environ)
    env["PYTHONPATH"] = str(ROOT / "src") + os.pathsep + env.get("PYTHONPATH", "")
    return subprocess.run(
        [sys.executable, *args],
        cwd=ROOT,
        env=env,
        text=True,
        capture_output=True,
        check=False,
    )


def _successful(*args: str) -> str:
    result = _call(*args)
    assert result.returncode == 0, result.stdout + "\n" + result.stderr
    return result.stdout


def _setup(tmp_path: Path) -> tuple[Path, Path, Path]:
    source = tmp_path / "original.xlsx"
    schema = tmp_path / "schema.yaml"
    challenge_dir = tmp_path / "challenge"
    _fixture(source)
    schema.write_text(
        "version: 1\n"
        "tables:\n"
        "  Cases:\n"
        "    grain_key: [case_id]\n"
        "  Diagnoses:\n"
        "    grain_key: [diagnosis_key]\n"
        "  Treatments:\n"
        "    grain_key: [treatment_key]\n"
        "  FollowUps:\n"
        "    grain_key: [followup_key]\n"
        "  LabTests:\n"
        "    grain_key: [labtest_key]\n"
        "foreign_keys:\n"
        "  - {child_table: Diagnoses, child_column: case_id, parent_table: Cases, parent_column: case_id}\n"
        "  - {child_table: Treatments, child_column: diagnosis_key, parent_table: Diagnoses, parent_column: diagnosis_key}\n"
        "  - {child_table: FollowUps, child_column: case_id, parent_table: Cases, parent_column: case_id}\n"
        "  - {child_table: LabTests, child_column: followup_key, parent_table: FollowUps, parent_column: followup_key}\n",
        encoding="utf-8",
    )
    return source, schema, challenge_dir


def test_seeded_challenge_plan_scoring_apply_and_undo(tmp_path: Path) -> None:
    source, schema, challenge_dir = _setup(tmp_path)
    pristine_hash = file_sha256(source)
    _successful(
        str(SCRIPT), "make", "--source", str(source),
        "--schema", str(schema), "--out", str(challenge_dir),
    )
    challenge = challenge_dir / "commpass_challenge.xlsx"
    answer_key = json.loads(
        (challenge_dir / "answer_key.json").read_text(encoding="utf-8")
    )
    assert len(answer_key["injections"]) == 4
    assert file_sha256(source) == pristine_hash
    original_book = load_workbook(source)
    challenge_book = load_workbook(challenge)
    assert len(original_book.sheetnames) == 5
    assert len(challenge_book.sheetnames) == 6
    assert "BenchmarkLabels" not in original_book.sheetnames
    original_book.close()
    challenge_book.close()

    plan_file = challenge_dir / "plan.json"
    _successful(
        "-m", "data_cleaning_assistant.cli", "plan", str(challenge),
        "--schema", str(challenge_dir / "study_schema.yaml"),
        "--output", str(plan_file),
    )
    plan = CleaningPlan.model_validate_json(plan_file.read_text(encoding="utf-8"))
    assert len(plan.actions) >= 4
    assert all(action.table != "BenchmarkLabels" for action in plan.actions)
    _successful(
        str(SCRIPT), "verify", "--challenge", str(challenge_dir),
        "--plan", str(plan_file),
    )
    _successful(
        str(SCRIPT), "verify", "--challenge", str(challenge_dir),
        "--plan", str(plan_file), "--exercise-apply-undo",
    )
    score = json.loads(
        (challenge_dir / "benchmark_score.json").read_text(encoding="utf-8")
    )
    assert score["expected_cells_covered"] == 4
    assert score["unexpected_action_targets"] == []
    assert score["synthetic_label_pairs_exposed_for_review"]
    assert score["ambiguous_assay_labels_not_in_deterministic_merges"]
    assert score["undo_exact_value_equality"]
    assert file_sha256(source) == pristine_hash
    assert file_sha256(challenge) == answer_key["challenge_sha256"]


def test_challenge_is_reproducible_and_never_overwrites(tmp_path: Path) -> None:
    source, schema, first = _setup(tmp_path)
    second = tmp_path / "second"
    for destination in (first, second):
        _successful(
            str(SCRIPT), "make", "--source", str(source),
            "--schema", str(schema), "--out", str(destination), "--seed", "37",
        )
    first_key = json.loads((first / "answer_key.json").read_text(encoding="utf-8"))
    second_key = json.loads((second / "answer_key.json").read_text(encoding="utf-8"))
    assert first_key["injections"] == second_key["injections"]
    again = _call(
        str(SCRIPT), "make", "--source", str(source),
        "--schema", str(schema), "--out", str(first),
    )
    assert again.returncode != 0
    assert "already exists" in again.stderr
