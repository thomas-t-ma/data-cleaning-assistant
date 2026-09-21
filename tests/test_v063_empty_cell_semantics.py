"""Real blank cells are coverage, not executable null-normalization defects."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from typer.testing import CliRunner

from data_cleaning_assistant.actions import ActionType
from data_cleaning_assistant.cli import app
from data_cleaning_assistant.engine import profile_dataset
from data_cleaning_assistant.io import load_dataset
from data_cleaning_assistant.planner import file_sha256, propose_actions
from data_cleaning_assistant.study_schema import StudySchema
from data_cleaning_assistant.transforms import preview_action


def _excel_with_blanks(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Study"
    sheet.append(["subject_id", "all_blank", "partial", "sentinel", "spaces"])
    sheet.append(["P001", None, None, "N/A", " "])
    sheet.append(["P002", None, "observed", "ok", "ok"])
    sheet.append(["P003", None, None, "ok", "ok"])
    book.save(path)


def test_excel_blanks_are_missing_coverage_not_hidden_missing_actions(
    tmp_path: Path,
) -> None:
    source = tmp_path / "study.xlsx"
    _excel_with_blanks(source)
    original_hash = file_sha256(source)
    loaded = load_dataset(source)
    # With na_filter=False pandas intentionally yields the same text for an
    # ordinary empty Excel cell and a literal zero-length text cell.
    assert loaded.tables["Study"]["all_blank"].tolist() == ["", "", ""]

    report = profile_dataset(source)
    profile = next(t for t in report.tables if t.name == "Study")
    columns = {c.name: c for c in profile.columns}
    assert columns["all_blank"].null_count == 3
    assert columns["all_blank"].null_percent == 100
    assert columns["all_blank"].sample_values == []
    assert columns["partial"].null_count == 2
    assert columns["partial"].null_percent == 2 / 3 * 100
    assert not [
        issue for issue in report.issues
        if issue.column in {"all_blank", "partial"}
        and issue.type == "hidden_missing_values"
    ]
    assert not [
        issue for issue in report.issues
        if issue.column == "all_blank" and issue.type == "constant_column"
    ]
    # One observed value among otherwise blank cells is still constant among
    # non-missing observations; that independent informational finding remains.

    plan = propose_actions(source)
    assert not [
        action for action in plan.actions
        if action.column in {"all_blank", "partial"}
    ]
    assert not [
        note for note in plan.planning_notes
        if note.column in {"all_blank", "partial"}
    ]
    assert file_sha256(source) == original_hash


def test_nonempty_sentinels_and_whitespace_remain_reviewable(tmp_path: Path) -> None:
    source = tmp_path / "study.xlsx"
    _excel_with_blanks(source)
    report = profile_dataset(source)
    findings = {
        issue.column: issue for issue in report.issues
        if issue.type == "hidden_missing_values"
    }
    assert findings["sentinel"].evidence["structural_tokens"] == {"N/A": 1}
    assert findings["spaces"].evidence["structural_tokens"] == {" ": 1}
    plan = propose_actions(source)
    missing = {
        action.column: action.parameters["values"]
        for action in plan.actions
        if action.type == ActionType.NORMALIZE_MISSING
    }
    assert missing == {"sentinel": ["N/A"], "spaces": [" "]}


def test_empty_csv_fields_are_coverage_and_nonempty_tokens_still_propose(
    tmp_path: Path,
) -> None:
    source = tmp_path / "study.csv"
    source.write_text(
        "subject_id,result,age\n"
        "P001,,N/A\n"
        "P002,7,44\n"
        "P003,,45\n",
        encoding="utf-8",
    )
    report = profile_dataset(source)
    columns = {c.name: c for c in report.tables[0].columns}
    assert columns["result"].null_count == 2
    assert not any(
        issue.column == "result" and issue.type == "hidden_missing_values"
        for issue in report.issues
    )
    plan = propose_actions(source)
    assert not any(action.column == "result" for action in plan.actions)
    assert any(
        action.column == "age"
        and action.type == ActionType.NORMALIZE_MISSING
        and action.parameters == {"values": ["N/A"]}
        for action in plan.actions
    )


def test_explicit_schema_can_still_select_empty_string(tmp_path: Path) -> None:
    source = tmp_path / "study.xlsx"
    _excel_with_blanks(source)
    schema = StudySchema.model_validate({
        "tables": {
            "Study": {
                "columns": {
                    "partial": {"missing_tokens": [""]}
                }
            }
        }
    })
    plan = propose_actions(source, schema=schema)
    actions = [
        action for action in plan.actions
        if action.column == "partial"
    ]
    assert len(actions) == 1
    assert actions[0].parameters == {"values": [""]}
    assert preview_action(load_dataset(source).tables, actions[0]).affected_count == 2


def test_cli_prints_blank_coverage_without_fictitious_actions(tmp_path: Path) -> None:
    source = tmp_path / "study.xlsx"
    _excel_with_blanks(source)
    runner = CliRunner()
    result = runner.invoke(app, ["profile", str(source)])
    assert result.exit_code == 0, result.output
    assert "Already-empty fields" in result.output
    plan_path = tmp_path / "plan.json"
    plan_result = runner.invoke(
        app, ["plan", str(source), "--output", str(plan_path)]
    )
    assert plan_result.exit_code == 0, plan_result.output
    from data_cleaning_assistant.actions import CleaningPlan

    plan = CleaningPlan.model_validate_json(plan_path.read_text(encoding="utf-8"))
    assert {a.column for a in plan.actions} == {"sentinel", "spaces"}
    assert not any(note.column == "all_blank" for note in plan.planning_notes)
