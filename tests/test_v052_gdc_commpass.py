"""Offline validation for the open GDC CoMMpass pilot loader."""

from __future__ import annotations

import importlib.util
import json
from pathlib import Path

from openpyxl import load_workbook

from data_cleaning_assistant.engine import profile_dataset
from data_cleaning_assistant.study_schema import load_study_schema

MODULE_PATH = Path(__file__).resolve().parents[1] / "examples" / "fetch_gdc_commpass.py"
spec = importlib.util.spec_from_file_location("fetch_gdc_commpass", MODULE_PATH)
assert spec is not None and spec.loader is not None
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def _case(case_id: str) -> dict:
    return {
        "case_id": case_id,
        "project": {"project_id": "MMRF-COMMPASS"},
        "disease_type": "Plasma Cell Tumors",
        "demographic": {"vital_status": "Alive"},
        "diagnoses": [
            {
                "days_to_last_follow_up": 120,
                "treatments": [
                    {
                        "therapeutic_agents": "Agent A",
                        "days_to_treatment_start": 1,
                    },
                    {
                        "therapeutic_agents": "Agent B",
                        "days_to_treatment_start": 45,
                    },
                ],
            }
        ],
        "follow_ups": [
            {
                "days_to_follow_up": 30,
                "disease_response": "Unknown",
                "molecular_tests": [
                    {
                        "test_analyte_type": "Creatinine",
                        "test_value": "1.2",
                        "test_units": "mg/dL",
                    },
                    {
                        "test_analyte_type": "Creatinine",
                        "test_value": "<5",
                        "test_units": "mg/dL",
                    },
                ],
            },
            {"days_to_follow_up": 90, "disease_response": "Stable Disease"},
        ],
    }


def test_fetch_pages_and_flatten_preserve_grain(monkeypatch) -> None:
    responses = iter([
        {"data": {"pagination": {"total": 3}, "hits": [_case("id-1"), _case("id-2")]}},
        {"data": {"pagination": {"total": 3}, "hits": [_case("id-3")]}},
    ])
    urls: list[str] = []

    def fake_get(url: str) -> dict:
        urls.append(url)
        return next(responses)

    monkeypatch.setattr(module, "get_json", fake_get)
    cases, request_urls, available = module.fetch_cases(limit=3, page_size=2)
    assert len(cases) == 3
    assert available == 3
    assert len(request_urls) == 2
    assert "from=2" in urls[1]
    assert "MMRF-COMMPASS" in urls[0]
    rows = module.linked_tables(cases)
    assert {name: len(rows[name]) for name in module.TABLES} == {
        "Cases": 3,
        "Diagnoses": 3,
        "Treatments": 6,
        "FollowUps": 6,
        "LabTests": 6,
    }
    assert rows["LabTests"][1]["test_value"] == "<5"
    assert rows["FollowUps"][0]["disease_response"] == "Unknown"
    assert all("submitter_id" not in row for table in rows.values() for row in table)
    assert len({r["labtest_key"] for r in rows["LabTests"]}) == 6
    assert len({r["followup_key"] for r in rows["FollowUps"]}) == 6


def test_project_guard_fails_closed(monkeypatch) -> None:
    monkeypatch.setattr(
        module,
        "get_json",
        lambda url: {
            "data": {
                "pagination": {"total": 1},
                "hits": [{"case_id": "out-of-cohort", "project": {"project_id": "OTHER"}}],
            }
        },
    )
    import pytest

    with pytest.raises(ValueError, match="outside the requested project"):
        module.fetch_cases(limit=1)


def test_workbook_and_schema_are_valid_on_synthetic_fixture(tmp_path: Path) -> None:
    rows = module.linked_tables([_case("id-1"), _case("id-2")])
    workbook = tmp_path / "pilot.xlsx"
    schema_path = tmp_path / "schema.yaml"
    module.write_workbook(workbook, rows)
    schema_path.write_text(module.build_schema(rows), encoding="utf-8")
    schema = load_study_schema(schema_path)
    report = profile_dataset(workbook, schema=schema)
    assert len(report.tables) == 5
    assert not any(
        issue.type == "schema_unmatched_foreign_key"
        for issue in report.issues
    )
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        assert list(book["LabTests"].values)[2][-1] is not None
    finally:
        book.close()
    assert json.loads(json.dumps(rows)) == rows
