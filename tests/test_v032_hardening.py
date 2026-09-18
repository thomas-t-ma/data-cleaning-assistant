from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from data_cleaning_assistant.actions import ActionType
from data_cleaning_assistant.io import load_dataset
from data_cleaning_assistant.planner import propose_actions
from data_cleaning_assistant.recipe import check_recipe_compatibility, load_recipe
from data_cleaning_assistant.transforms import (
    apply_plan,
    undo_changeset,
    write_cleaned_dataset,
    write_recipe,
    write_tables,
)


def test_exact_duplicate_removal_is_reversible_after_reload(tmp_path: Path) -> None:
    source = tmp_path / "duplicates.csv"
    source.write_text(
        "patient_id,status\n"
        "P001,Active\n"
        "P001,Active\n"
        "P002,Inactive\n",
        encoding="utf-8",
    )

    loaded = load_dataset(source)
    original = loaded.tables[source.stem].copy(deep=True)
    plan = propose_actions(source)
    duplicate = next(
        action
        for action in plan.actions
        if action.type == ActionType.REMOVE_EXACT_DUPLICATES
    )

    cleaned, changeset = apply_plan(
        source,
        loaded.tables,
        plan,
        approved_ids={duplicate.id},
    )

    assert len(cleaned[source.stem]) == 2
    assert changeset.affected_row_count == 1
    assert changeset.affected_cell_count == 0

    cleaned_path = write_tables(cleaned, tmp_path / "cleaned.csv")
    reloaded = load_dataset(cleaned_path)
    restored = undo_changeset(reloaded.tables, changeset)
    restored_frame = next(iter(restored.values()))

    pd.testing.assert_frame_equal(restored_frame, original)


def test_surrounding_whitespace_gets_reviewable_trim_action(tmp_path: Path) -> None:
    source = tmp_path / "whitespace.csv"
    source.write_text(
        "patient_id,status\n"
        "P001, Active \n"
        "P002,Inactive\n",
        encoding="utf-8",
    )

    loaded = load_dataset(source)
    plan = propose_actions(source)
    trim = next(
        action
        for action in plan.actions
        if action.type == ActionType.TRIM_WHITESPACE and action.column == "status"
    )

    cleaned, changeset = apply_plan(
        source,
        loaded.tables,
        plan,
        approved_ids={trim.id},
    )

    assert cleaned[source.stem].loc[0, "status"] == "Active"
    assert changeset.affected_cell_count == 1


def _make_styled_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.freeze_panes = "A2"
    ws.column_dimensions["B"].width = 24
    ws.append(["PatientID", "Status", "Formula"])
    ws.append(["P001", " Active ", "=1+1"])
    ws.append(["P002", "Inactive", "=2+2"])
    ws["B2"].font = Font(bold=True)
    ws["B2"].fill = PatternFill(fill_type="solid", fgColor="00FF00")
    wb.save(path)


def test_excel_cleaning_preserves_styles_layout_and_untouched_formulas(tmp_path: Path) -> None:
    source = tmp_path / "styled.xlsx"
    output = tmp_path / "styled_cleaned.xlsx"
    _make_styled_workbook(source)

    loaded = load_dataset(source)
    plan = propose_actions(source)
    trim = next(
        action
        for action in plan.actions
        if action.type == ActionType.TRIM_WHITESPACE
        and action.table == "Data"
        and action.column == "Status"
    )
    cleaned, _ = apply_plan(
        source,
        loaded.tables,
        plan,
        approved_ids={trim.id},
    )

    write_cleaned_dataset(source, loaded.tables, cleaned, output)

    wb = load_workbook(output, data_only=False)
    ws = wb["Data"]
    assert ws["B2"].value == "Active"
    assert ws["B2"].font.bold is True
    assert ws["B2"].fill.fgColor.rgb == "0000FF00"
    assert ws["C2"].value == "=1+1"
    assert ws.column_dimensions["B"].width == 24
    assert ws.freeze_panes == "A2"


def test_excel_row_deletion_refuses_formula_reference_risk(tmp_path: Path) -> None:
    source = tmp_path / "formula_rows.xlsx"
    output = tmp_path / "cleaned.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["A", "B"])
    ws.append([1, "=A2"])
    ws.append([1, "=A2"])
    wb.save(source)

    original = {"Data": pd.DataFrame({"A": [1, 1], "B": ["=A2", "=A2"]})}
    cleaned = {"Data": original["Data"].iloc[[0]].copy()}

    with pytest.raises(ValueError, match="contains formulas"):
        write_cleaned_dataset(source, original, cleaned, output)


def test_write_cleaned_dataset_refuses_in_place_overwrite(tmp_path: Path) -> None:
    source = tmp_path / "source.csv"
    source.write_text("a\n1\n", encoding="utf-8")
    loaded = load_dataset(source)

    with pytest.raises(ValueError, match="overwrite"):
        write_cleaned_dataset(source, loaded.tables, loaded.tables, source)


def test_recipe_reports_nonblocking_schema_drift(tmp_path: Path) -> None:
    source = tmp_path / "wave1.csv"
    target = tmp_path / "wave2.csv"
    recipe_path = tmp_path / "recipe.yaml"

    source.write_text(
        "patient_id,sex,age,status\n"
        "P001,Female,30, Active \n"
        "P002,Male,40,Inactive\n"
        "P003,Female,50,Active\n",
        encoding="utf-8",
    )
    target.write_text(
        "patient_id,sex,age,status,site\n"
        "P101,Female,20,Active,A\n"
        "P102,Other,40,N/A,A\n"
        "P103,Male,90,Active,B\n",
        encoding="utf-8",
    )

    loaded = load_dataset(source)
    plan = propose_actions(source)
    trim = next(
        action
        for action in plan.actions
        if action.type == ActionType.TRIM_WHITESPACE and action.column == "status"
    )
    _, changeset = apply_plan(
        source,
        loaded.tables,
        plan,
        approved_ids={trim.id},
    )
    write_recipe(plan, changeset, recipe_path)

    recipe = load_recipe(recipe_path)
    target_tables = load_dataset(target).tables
    report = check_recipe_compatibility(recipe, target_tables)

    assert report.compatible
    kinds = {item.kind for item in report.drift}
    assert "new_column" in kinds
    assert "new_categories" in kinds
    assert "range_expansion" in kinds
    assert "missingness_change" in kinds


def test_excel_duplicate_removal_can_be_undone_with_row_reinsertion(tmp_path: Path) -> None:
    source = tmp_path / "styled_duplicates.xlsx"
    cleaned_path = tmp_path / "styled_duplicates_cleaned.xlsx"
    restored_path = tmp_path / "styled_duplicates_restored.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.column_dimensions["B"].width = 22
    ws.append(["PatientID", "Status"])
    ws.append(["P001", "Active"])
    ws.append(["P001", "Active"])
    ws.append(["P002", "Inactive"])
    ws["B2"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    ws["B3"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    wb.save(source)

    loaded = load_dataset(source)
    original_frame = loaded.tables["Data"].copy(deep=True)
    plan = propose_actions(source)
    duplicate = next(
        action
        for action in plan.actions
        if action.type == ActionType.REMOVE_EXACT_DUPLICATES
    )
    cleaned, changeset = apply_plan(
        source,
        loaded.tables,
        plan,
        approved_ids={duplicate.id},
    )
    write_cleaned_dataset(source, loaded.tables, cleaned, cleaned_path)

    reloaded = load_dataset(cleaned_path)
    restored = undo_changeset(reloaded.tables, changeset)
    inserted_rows = {
        "Data": [
            row
            for execution in changeset.executions
            for row in execution.deleted_rows
            if execution.table == "Data"
        ]
    }
    write_cleaned_dataset(
        cleaned_path,
        reloaded.tables,
        restored,
        restored_path,
        inserted_rows=inserted_rows,
    )

    final_loaded = load_dataset(restored_path)
    pd.testing.assert_frame_equal(final_loaded.tables["Data"], original_frame)

    final_wb = load_workbook(restored_path)
    final_ws = final_wb["Data"]
    assert final_ws.column_dimensions["B"].width == 22
    assert final_ws["B2"].fill.fgColor.rgb == "00FFFF00"
    assert final_ws["B3"].fill.fgColor.rgb == "00FFFF00"
