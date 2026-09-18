from __future__ import annotations

from pathlib import Path

from data_cleaning_assistant import profile_dataset


def test_profile_detects_core_issues(tmp_path: Path) -> None:
    csv_path = tmp_path / "messy.csv"
    csv_path.write_text(
        "patient_id,sex,age,visit_date,site\n"
        "P001,M,35,2026-01-01,Atlanta\n"
        "P002,m,41,01/03/2026,Atlanta\n"
        "P003, M ,unknown,not-a-date,Atlanta\n"
        "P003, M ,unknown,not-a-date,Atlanta\n"
        "P004,M,forty,2026-01-05,Atlanta\n"
        "P005,M,50,2026-01-06,Atlanta\n",
        encoding="utf-8",
    )

    report = profile_dataset(csv_path)
    issue_types = {issue.type for issue in report.issues}

    assert report.tables[0].row_count == 6
    assert "potential_identifier" in issue_types
    assert "category_inconsistency" in issue_types
    assert "hidden_missing_values" in issue_types
    assert "mixed_numeric_values" in issue_types
    assert "malformed_or_mixed_dates" in issue_types
    assert "exact_duplicate_rows" in issue_types
    assert "constant_column" in issue_types


def test_source_is_not_modified(tmp_path: Path) -> None:
    csv_path = tmp_path / "source.csv"
    contents = "a,b\n1,x\n2,y\n"
    csv_path.write_text(contents, encoding="utf-8")
    profile_dataset(csv_path)
    assert csv_path.read_text(encoding="utf-8") == contents


def test_hidden_missing_token_is_not_type_error(tmp_path: Path) -> None:
    csv_path = tmp_path / "missing_numeric.csv"
    csv_path.write_text("age\n10\n11\nN/A\n12\n", encoding="utf-8")

    report = profile_dataset(csv_path)
    age_issues = [issue for issue in report.issues if issue.column == "age"]
    age_types = {issue.type for issue in age_issues}

    assert "hidden_missing_values" in age_types
    assert "mixed_numeric_values" not in age_types


def test_excel_multisheet_and_json_serialization(tmp_path: Path) -> None:
    from openpyxl import Workbook

    workbook_path = tmp_path / "book.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Patients"
    ws1.append(["PatientID", "Age"])
    ws1.append(["P1", 30])
    ws2 = wb.create_sheet("Visits")
    ws2.append(["PatientID", "VisitDate"])
    ws2.append(["P1", "2026-01-01"])
    wb.save(workbook_path)

    report = profile_dataset(workbook_path)

    assert {table.name for table in report.tables} == {"Patients", "Visits"}
    assert '"source_type":"excel"' in report.model_dump_json()
