from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from pydantic import ValidationError
from typer.testing import CliRunner

from data_cleaning_assistant.cli import app
from data_cleaning_assistant.schema_draft import (
    SchemaSuggestionDecision,
    SchemaSuggestionKind,
    build_schema_draft,
    load_schema_draft,
    materialize_study_schema,
    write_schema_draft,
)
from data_cleaning_assistant.study_schema import load_study_schema


def _write_single_table(path: Path) -> None:
    path.write_text(
        "SubjectID,Status,Age\n"
        "S001,Active,34\n"
        "S002,active,NA\n"
        "S003,Active,40\n",
        encoding="utf-8",
    )


def _write_related_workbook(path: Path) -> None:
    workbook = Workbook()
    subjects = workbook.active
    subjects.title = "Subjects"
    subjects.append(["SubjectID", "Status"])
    subjects.append(["S001", "Active"])
    subjects.append(["S002", "Active"])
    subjects.append(["S003", "Inactive"])

    visits = workbook.create_sheet("Visits")
    visits.append(["SubjectID", "Visit", "Result"])
    visits.append(["S001", "1", "10"])
    visits.append(["S001", "2", "11"])
    visits.append(["S002", "1", "12"])
    visits.append(["S003", "1", "13"])
    workbook.save(path)


def test_schema_draft_uses_existing_structural_and_cleaning_evidence(
    tmp_path: Path,
) -> None:
    source = tmp_path / "study.csv"
    _write_single_table(source)

    draft = build_schema_draft(source)

    kinds = [suggestion.kind for suggestion in draft.suggestions]
    assert SchemaSuggestionKind.GRAIN_KEY in kinds
    assert SchemaSuggestionKind.CANONICAL_MAPPING in kinds
    assert SchemaSuggestionKind.MISSING_TOKEN in kinds

    grain = next(
        item
        for item in draft.suggestions
        if item.kind == SchemaSuggestionKind.GRAIN_KEY
    )
    assert grain.proposed == {"columns": ["SubjectID"]}
    assert grain.confidence >= 0.95

    canonical = next(
        item
        for item in draft.suggestions
        if item.kind == SchemaSuggestionKind.CANONICAL_MAPPING
    )
    assert canonical.table == "study"
    assert canonical.column == "Status"
    assert canonical.proposed == {"source": "active", "target": "Active"}

    missing = next(
        item
        for item in draft.suggestions
        if item.kind == SchemaSuggestionKind.MISSING_TOKEN
    )
    assert missing.table == "study"
    assert missing.column == "Age"
    assert missing.proposed == {"token": "NA"}


def test_schema_draft_is_not_loadable_as_authoritative_schema(
    tmp_path: Path,
) -> None:
    source = tmp_path / "study.csv"
    draft_path = tmp_path / "draft.yaml"
    _write_single_table(source)

    write_schema_draft(build_schema_draft(source), draft_path)

    with pytest.raises(ValidationError):
        load_study_schema(draft_path)


def test_only_explicitly_approved_draft_items_materialize(
    tmp_path: Path,
) -> None:
    source = tmp_path / "study.csv"
    _write_single_table(source)
    draft = build_schema_draft(source)

    for suggestion in draft.suggestions:
        if suggestion.kind in {
            SchemaSuggestionKind.GRAIN_KEY,
            SchemaSuggestionKind.CANONICAL_MAPPING,
        }:
            suggestion.decision = SchemaSuggestionDecision.APPROVED
        else:
            suggestion.decision = SchemaSuggestionDecision.REJECTED

    schema = materialize_study_schema(
        draft,
        dataset_path=source,
    )

    table = schema.tables["study"]
    assert table.grain_key == ["SubjectID"]
    assert table.columns["Status"].canonical_map == {"active": "Active"}
    assert "Age" not in table.columns


def test_schema_draft_proposes_high_confidence_many_to_one_foreign_key(
    tmp_path: Path,
) -> None:
    source = tmp_path / "study.xlsx"
    _write_related_workbook(source)

    draft = build_schema_draft(source)

    foreign_keys = [
        item
        for item in draft.suggestions
        if item.kind == SchemaSuggestionKind.FOREIGN_KEY
    ]
    assert len(foreign_keys) == 1
    suggestion = foreign_keys[0]
    assert suggestion.confidence >= 0.90
    assert suggestion.proposed == {
        "child_table": "Visits",
        "child_column": "SubjectID",
        "parent_table": "Subjects",
        "parent_column": "SubjectID",
    }


def test_schema_draft_and_review_cli_produce_confirmed_schema(
    tmp_path: Path,
) -> None:
    source = tmp_path / "study.csv"
    draft_path = tmp_path / "draft.yaml"
    schema_path = tmp_path / "study_schema.yaml"
    _write_single_table(source)
    runner = CliRunner()

    draft_result = runner.invoke(
        app,
        [
            "schema-draft",
            str(source),
            "--output",
            str(draft_path),
        ],
    )
    assert draft_result.exit_code == 0, draft_result.output
    assert "Nothing in this draft is authoritative" in draft_result.output

    draft = load_schema_draft(draft_path)
    assert len(draft.suggestions) == 3

    review_result = runner.invoke(
        app,
        [
            "schema-review",
            str(source),
            str(draft_path),
            "--output",
            str(schema_path),
        ],
        input="a\na\na\n",
    )
    assert review_result.exit_code == 0, review_result.output

    schema = load_study_schema(schema_path)
    assert schema.tables["study"].grain_key == ["SubjectID"]
    assert schema.tables["study"].columns["Status"].canonical_map == {
        "active": "Active"
    }
    assert schema.tables["study"].columns["Age"].missing_tokens == ["NA"]

    reviewed = load_schema_draft(
        schema_path.with_name("study_schema.reviewed-draft.yaml")
    )
    assert all(
        item.decision == SchemaSuggestionDecision.APPROVED
        for item in reviewed.suggestions
    )


def test_schema_review_can_backtrack_without_authorizing_previous_choice(
    tmp_path: Path,
) -> None:
    source = tmp_path / "study.csv"
    draft_path = tmp_path / "draft.yaml"
    schema_path = tmp_path / "study_schema.yaml"
    _write_single_table(source)
    draft = build_schema_draft(source)
    write_schema_draft(draft, draft_path)

    result = CliRunner().invoke(
        app,
        [
            "schema-review",
            str(source),
            str(draft_path),
            "--output",
            str(schema_path),
        ],
        input="a\nb\nr\nr\nr\n",
    )

    assert result.exit_code == 0, result.output
    assert "Previous schema decision reset" in result.output

    reviewed = load_schema_draft(
        schema_path.with_name("study_schema.reviewed-draft.yaml")
    )
    assert all(
        item.decision == SchemaSuggestionDecision.REJECTED
        for item in reviewed.suggestions
    )

    schema = load_study_schema(schema_path)
    assert schema.tables == {}
    assert schema.foreign_keys == []
