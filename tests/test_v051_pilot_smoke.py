"""End-to-end CLI smoke test using synthetic research data only."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from typer.testing import CliRunner

from data_cleaning_assistant.actions import ActionDecision, ActionType, CleaningPlan
from data_cleaning_assistant.cli import app
from data_cleaning_assistant.planner import file_sha256


EXAMPLE_SCHEMA = (
    Path(__file__).resolve().parents[1] / "examples" / "study_schema_example.yaml"
)


def _synthetic_workbook(path: Path) -> None:
    wb = Workbook()
    visits = wb.active
    visits.title = "Visits"
    visits.append(
        ["PatientID", "Sex", "Age", "VisitDate", "Site", "Hemoglobin", "StudyArm"]
    )
    for row in (
        ("P001", "Female", "34", "2026-01-02", "Atlanta", "11.2", "A"),
        ("P002", "female", "41", "1/5/2026", "Atlanta", "12.1", "A"),
        ("P003", " Female ", "N/A", "January 8, 2026", "Atlanta", "10.8", "B"),
        ("P004", "Male", "52", "bad-date", "Atlanta", "42.0", "B"),
        ("P005", "MALE", "49", "2026-01-12", "Atlanta", "unknown", "B"),
        ("P005", "MALE", "49", "2026-01-12", "Atlanta", "unknown", "B"),
    ):
        visits.append(row)

    subjects = wb.create_sheet("Subjects")
    subjects.append(["PatientID", "DOB", "Status"])
    for row in (
        ("P001", "1991-04-02", "Active"),
        ("P002", "1984-03-11", "active"),
        ("P003", "NA", "Active"),
    ):
        subjects.append(row)
    wb.save(path)


def _cells(path: Path) -> dict[str, list[tuple[object, ...]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            sheet.title: list(sheet.values)
            for sheet in wb.worksheets
        }
    finally:
        wb.close()


def test_synthetic_workbook_complete_cli_pilot(tmp_path: Path) -> None:
    """Profile, plan, preview, review, apply, replay preview, and undo."""
    source = tmp_path / "synthetic_study.xlsx"
    profile_path = tmp_path / "profile.json"
    plan_path = tmp_path / "plan.json"
    reviewed_path = tmp_path / "reviewed.json"
    cleaned_path = tmp_path / "cleaned.xlsx"
    restored_path = tmp_path / "restored.xlsx"
    _synthetic_workbook(source)
    original_hash = file_sha256(source)
    original_cells = _cells(source)
    runner = CliRunner()

    def invoke(*args: str, input_text: str | None = None) -> str:
        result = runner.invoke(app, list(args), input=input_text)
        assert result.exit_code == 0, result.output
        return result.output

    invoke(
        "profile", str(source), "--schema", str(EXAMPLE_SCHEMA),
        "--output", str(profile_path),
    )
    report = json.loads(profile_path.read_text(encoding="utf-8"))
    types = {issue["type"] for issue in report["issues"]}
    assert "schema_unmatched_foreign_key" in types
    assert "schema_duplicate_grain" in types

    invoke(
        "plan", str(source), "--schema", str(EXAMPLE_SCHEMA),
        "--output", str(plan_path),
    )
    plan = CleaningPlan.model_validate_json(plan_path.read_text(encoding="utf-8"))
    assert plan.study_schema_sha256
    assert len(plan.actions) >= 2
    assert all(a.decision == ActionDecision.PROPOSED for a in plan.actions)

    invoke("preview", str(source), str(plan_path))

    # Never accept row deletion in the pilot; approve only exact dictionary
    # category mappings so restoration is straightforward to inspect.
    approved = [
        action for action in plan.actions
        if action.type == ActionType.NORMALIZE_CATEGORIES
        and action.column in {"Sex", "Status"}
    ]
    assert len(approved) == 2
    answers = [
        "a" if action in approved else (
            "r" if action.type == ActionType.REMOVE_EXACT_DUPLICATES else "s"
        )
        for action in plan.actions
    ]
    invoke(
        "review", str(source), str(plan_path),
        "--output", str(reviewed_path),
        input_text="\\n".join(answers) + "\\n",
    )
    reviewed = CleaningPlan.model_validate_json(
        reviewed_path.read_text(encoding="utf-8")
    )
    assert sum(a.decision == ActionDecision.APPROVED for a in reviewed.actions) == 2
    assert all(
        a.decision == ActionDecision.REJECTED
        for a in reviewed.actions
        if a.type == ActionType.REMOVE_EXACT_DUPLICATES
    )

    invoke(
        "apply", str(source), str(reviewed_path),
        "--output", str(cleaned_path),
    )
    assert file_sha256(source) == original_hash
    assert cleaned_path.exists()
    assert _cells(cleaned_path)["Visits"][2][1] == "Female"
    assert _cells(cleaned_path)["Subjects"][2][2] == "Active"
    # Both intentionally repeated measurements survive.
    assert len(_cells(cleaned_path)["Visits"]) == len(original_cells["Visits"])

    audit_path = cleaned_path.with_name(f"{cleaned_path.name}.audit.json")
    recipe_path = cleaned_path.with_name(f"{cleaned_path.name}.recipe.yaml")
    audit = json.loads(audit_path.read_text(encoding="utf-8"))
    assert audit["study_schema_sha256"] == plan.study_schema_sha256
    assert len(audit["executions"]) == 2
    assert recipe_path.exists()

    replay_output = invoke("replay", str(source), str(recipe_path))
    assert "Preview only" in replay_output

    invoke(
        "undo", str(cleaned_path), str(audit_path),
        "--output", str(restored_path),
    )
    assert _cells(restored_path) == original_cells
    assert file_sha256(source) == original_hash
