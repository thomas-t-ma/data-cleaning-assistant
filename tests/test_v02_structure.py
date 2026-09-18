from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from data_cleaning_assistant import profile_dataset


def _make_research_book(path: Path) -> None:
    wb = Workbook()
    patients = wb.active
    patients.title = "Patients"
    patients.append(["PatientID", "Sex", "Diagnosis"])
    patients.append(["P001", "Female", "Myeloma"])
    patients.append(["P002", "Male", "Myeloma"])
    patients.append(["P003", "Female", "Control"])

    visits = wb.create_sheet("Visits")
    visits.append(["PatientID", "VisitDate", "VisitType", "Site"])
    visits.append(["P001", "2026-01-01", "Baseline", "Atlanta"])
    visits.append(["P001", "2026-02-01", "Follow-up", "Atlanta"])
    visits.append(["P002", "2026-01-04", "Baseline", "Atlanta"])
    visits.append(["P003", "2026-01-05", "Basline", "Atlanta"])
    wb.save(path)


def test_candidate_keys_grain_and_relationships(tmp_path: Path) -> None:
    path = tmp_path / "research.xlsx"
    _make_research_book(path)
    report = profile_dataset(path)

    profiles = {profile.name: profile for profile in report.tables}
    assert profiles["Patients"].grain is not None
    assert profiles["Patients"].grain.label == "subject"
    assert any(key.columns == ["PatientID"] for key in profiles["Patients"].candidate_keys)

    assert profiles["Visits"].grain is not None
    assert profiles["Visits"].grain.label in {"visit", "observation"}
    assert any(
        rel.left_column == "PatientID" and rel.right_column == "PatientID"
        for rel in report.relationships
    )
    assert any(issue.type == "repeated_entity_rows" for issue in report.issues)


def test_fuzzy_category_suggestions_are_review_only(tmp_path: Path) -> None:
    path = tmp_path / "categories.csv"
    path.write_text(
        "subject_id,visit_type\n"
        "S1,Baseline\n"
        "S2,Basline\n"
        "S3,Follow-up\n"
        "S4,Followup\n",
        encoding="utf-8",
    )
    report = profile_dataset(path)
    fuzzy = [issue for issue in report.issues if issue.type == "fuzzy_category_match"]
    assert fuzzy
    assert all(issue.auto_fix_safe is False for issue in fuzzy)
